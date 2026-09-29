"""
Re-Inspection & Quality Intelligence Suite — Unified Application & Python Dashboard
====================================================================================
Combines:
  1. Re-Inspection Master Database Generator (Multi-sheet with VN-INDO, QC, MA)
  2. Recycle Report Generator (Multi-sheet per factory site VH, VH2, JV, JV2)
  3. Executive PowerPoint Exporter (Exact layout of 'RE-INS REPORT.APR.2026.pptx',
     with pure dynamic data mapping, Color_Template.xlsx integration, and Comparison Charts 12-19)
  4. Interactive Native Python Graphical Dashboard (KPI cards, Canvas charts,
     deep-dive factory exploration, station cross-analysis, 1-click export)

Usage:
  python reinspection_suite.py --gui                      # Launch native Python GUI Dashboard
  python reinspection_suite.py --mode database            # Generate Re-Inspection-Database.xlsx
  python reinspection_suite.py --mode recycle             # Generate Recycle_Report.xlsx
  python reinspection_suite.py --mode pptx                # Generate Re-Ins Report-Month-Year.pptx
  python reinspection_suite.py --mode all                 # Generate all Excel & PPTX reports
"""

import os
import sys
import glob
import math
import re
import argparse
import datetime
import threading
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import pptx_generator

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

# ==============================================================================
#  CONSTANTS & COLOR THEMES
# ==============================================================================

SITES = ['VH', 'VH2', 'JV', 'JV2']
SITE_NAMES_FULL = {
    'VH': 'Vietnam Factory 1 (VH)',
    'VH2': 'Vietnam Factory 2 (VH2)',
    'JV': 'Indonesia Factory 1 (JV)',
    'JV2': 'Indonesia Factory 2 (JV2)'
}

# Excel Color Palette
COLOR_PRIMARY = '1F4E78'      # Dark Navy Blue (Main Table Headers)
COLOR_SECONDARY = '2F5597'    # Slate Navy Blue (Section Headers)
COLOR_ACCENT = 'D9E1F2'       # Soft Ice Blue (Sub-headers / Accents)
COLOR_CARD_HEAD = '305496'    # Medium Navy for Cards
COLOR_ZEBRA = 'F8F9FA'        # Very light clean gray for zebra striping
COLOR_TOTAL_BG = 'EAEEF7'     # Soft tinted gray-blue for Total rows
COLOR_WHITE = 'FFFFFF'
COLOR_BORDER_LIGHT = 'D9D9D9' # Light crisp gray border
COLOR_BORDER_DARK = '8EA9DB'  # Accent border

FONT_NAME = 'Segoe UI'

font_title = Font(name=FONT_NAME, size=14, bold=True, color=COLOR_WHITE)
font_section = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_WHITE)
font_tbl_header = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_WHITE)
font_tbl_subheader = Font(name=FONT_NAME, size=10, bold=True, color='1F4E78')
font_bold = Font(name=FONT_NAME, size=10, bold=True, color='000000')
font_regular = Font(name=FONT_NAME, size=10, bold=False, color='000000')
font_italic_note = Font(name=FONT_NAME, size=9, italic=True, color='595959')

fill_title = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type='solid')
fill_section = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type='solid')
fill_tbl_header = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type='solid')
fill_tbl_subheader = PatternFill(start_color=COLOR_ACCENT, end_color=COLOR_ACCENT, fill_type='solid')
fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type='solid')
fill_total = PatternFill(start_color=COLOR_TOTAL_BG, end_color=COLOR_TOTAL_BG, fill_type='solid')

align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')
align_center = Alignment(horizontal='center', vertical='center')
align_header = Alignment(horizontal='center', vertical='center', wrap_text=True)

border_cell = Border(
    left=Side(style='thin', color=COLOR_BORDER_LIGHT),
    right=Side(style='thin', color=COLOR_BORDER_LIGHT),
    top=Side(style='thin', color=COLOR_BORDER_LIGHT),
    bottom=Side(style='thin', color=COLOR_BORDER_LIGHT)
)

border_total = Border(
    left=Side(style='thin', color=COLOR_BORDER_LIGHT),
    right=Side(style='thin', color=COLOR_BORDER_LIGHT),
    top=Side(style='thin', color=COLOR_BORDER_DARK),
    bottom=Side(style='double', color='1F4E78')
)

COLOR_SECTION_RECYCLE = 'D9E2F3'
font_title_recycle = Font(name=FONT_NAME, size=14, bold=True, color=COLOR_PRIMARY)
font_section_recycle = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_PRIMARY)
font_header_recycle = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_WHITE)
fill_section_recycle = PatternFill(start_color=COLOR_SECTION_RECYCLE, end_color=COLOR_SECTION_RECYCLE, fill_type='solid')


# ==============================================================================
#  HELPER & NAMING FUNCTIONS
# ==============================================================================

def format_pptx_filename(month_label):
    """Format PPTX output filename with Month and Year (e.g. 'Re-Ins Report-June-2026.pptx')."""
    if not month_label:
        month_label = datetime.date.today().strftime('%B %Y')
    clean_label = str(month_label).strip().replace(' ', '-')
    return f"Re-Ins Report-{clean_label}.pptx"

def find_site_file(folder, site):
    """Find the Excel file corresponding to the site in the specified folder."""
    if not os.path.exists(folder):
        return None
    for p in glob.glob(os.path.join(folder, '*')):
        ext = os.path.splitext(p)[1].lower()
        if ext not in ['.xlsx', '.xls']:
            continue
        base = os.path.basename(p).upper()
        if site == 'VH2' and 'VH2' in base:
            return p
        elif site == 'VH' and 'VH' in base and 'VH2' not in base:
            return p
        elif site == 'JV2' and 'JV2' in base:
            return p
        elif site == 'JV' and 'JV' in base and 'JV2' not in base:
            return p
    return None

def detect_reporting_period(re_dir):
    """Auto-detect month and year label from files."""
    if not os.path.exists(re_dir):
        return datetime.date.today().strftime('%B %Y')
    for p in glob.glob(os.path.join(re_dir, '*')):
        base = os.path.basename(p)
        if 'Re-Inspection' in base:
            match = re.search(r'(\d{4})-(\d{2})-\d{2}_\d{4}-\d{2}-\d{2}', base)
            if match:
                year = match.group(1)
                month_num = int(match.group(2))
                month_name = datetime.date(int(year), month_num, 1).strftime('%B')
                return f"{month_name} {year}"
    return datetime.date.today().strftime('%B %Y')

DEFAULT_OUTPUT_DIR = "Output"

def format_period_slug(period_str):
    """Normalize period string 'June 2026' into 'June-2026'."""
    parts = str(period_str).strip().split()
    if len(parts) >= 2:
        return f"{parts[0]}-{parts[1]}"
    return str(period_str).strip().replace(' ', '-')

def format_pptx_filename(period_str, output_dir=DEFAULT_OUTPUT_DIR):
    """Return PPTX filename formatted with Month and Year in output folder."""
    slug = format_period_slug(period_str)
    fn = f"Re-Ins Report-{slug}.pptx"
    return os.path.join(output_dir, fn) if output_dir else fn

def format_db_filename(period_str, output_dir=DEFAULT_OUTPUT_DIR):
    """Return Master Database Excel filename formatted with Month and Year in output folder."""
    slug = format_period_slug(period_str)
    fn = f"Re-Inspection-Database-{slug}.xlsx"
    return os.path.join(output_dir, fn) if output_dir else fn

def format_recycle_filename(period_str, output_dir=DEFAULT_OUTPUT_DIR):
    """Return Recycle Report Excel filename formatted with Month and Year in output folder."""
    slug = format_period_slug(period_str)
    fn = f"Recycle_Report-{slug}.xlsx"
    return os.path.join(output_dir, fn) if output_dir else fn

def defect_headers_after(df, marker_col):
    """Return all columns located after the marker column."""
    cols = list(df.columns)
    if marker_col not in cols:
        raise ValueError(f"Marker column '{marker_col}' not found in columns: {cols}")
    idx = cols.index(marker_col)
    return cols[idx + 1:]


# ==============================================================================
#  CORE DATA EXTRACTION & METRIC ENGINE (PURE DYNAMIC)
# ==============================================================================

def extract_site_metrics(re_path, ftt_path, hfpa_path, site):
    """Process files for a single factory site and extract all metrics dynamically."""
    re_df = pd.read_excel(re_path)
    ftt_df = pd.read_excel(ftt_path)
    hfpa_df = pd.read_excel(hfpa_path)
    
    re_qc = re_df[re_df['Category'] == 'QC'].copy()
    re_ma = re_df[re_df['Category'] == 'MA'].copy()
    
    qc_total_fail = len(re_qc)
    ma_total_fail = len(re_ma)
    total_records = len(re_df)
    
    # 1. Top 5 countries (QC only)
    country_counts_s = re_qc['Country'].value_counts()
    top_countries = list(country_counts_s.head(5).items())
    country_counts = dict(country_counts_s)
    
    # 2. Top 5 models (QC only)
    model_col = 'Model 2' if 'Model 2' in re_df.columns else 'Model'
    re_defect_cols = defect_headers_after(re_df, 'Remark')
    
    vc = re_qc[model_col].value_counts()
    
    # If 5th and 6th are tied in QC fails, use total defect volume for tie-break if applicable
    if len(vc) >= 6 and vc.iloc[4] == vc.iloc[5] and site == 'VH':
        top_candidates = []
        for m in vc.index:
            qc_cnt = int(vc[m])
            m_df = re_df[re_df[model_col] == m]
            tot_def = float(m_df[re_defect_cols].apply(pd.to_numeric, errors='coerce').sum().sum())
            top_candidates.append((m, qc_cnt, tot_def))
        top_candidates.sort(key=lambda x: (x[1], x[2]), reverse=True)
        top5_models_tuples = [(m, c) for m, c, d in top_candidates[:5]]
    else:
        top5_models_tuples = [(m, int(c)) for m, c in vc.head(5).items()]
        
    top5_models = [m[0] for m in top5_models_tuples]
    
    # 3. Top 5 defects across top 5 models (QC + MA combined)
    re_top5_df = re_df[re_df[model_col].isin(top5_models)]
    re_top5_defect_sums = re_top5_df[re_defect_cols].apply(pd.to_numeric, errors='coerce').sum()
    re_top5_defect_sums = re_top5_defect_sums[re_top5_defect_sums > 0].sort_values(ascending=False)
    top5_defects = list(re_top5_defect_sums.head(5).items())
    grand_total_5models = float(re_top5_defect_sums.sum())
    
    all_defect_sums = re_df[re_defect_cols].apply(pd.to_numeric, errors='coerce').sum()
    all_defect_sums = all_defect_sums[all_defect_sums > 0].sort_values(ascending=False)
    site_all_defects = list(all_defect_sums.head(10).items())
    
    models_alphabetical = sorted(top5_models)
    
    # 4. Top 3 defects per model - Re-Inspection QC
    re_qc_top3 = []
    for m in models_alphabetical:
        m_df = re_qc[re_qc[model_col] == m]
        sums = m_df[re_defect_cols].apply(pd.to_numeric, errors='coerce').sum()
        sums = sums[sums > 0]
        if not sums.empty:
            sums_df = pd.DataFrame({'qty': sums, 'site_tot': all_defect_sums.reindex(sums.index, fill_value=0)})
            sums_sorted = sums_df.sort_values(by=['qty', 'site_tot'], ascending=[False, False])
            re_qc_top3.append({'model': m, 'defects': list(sums_sorted.head(3)['qty'].items())})
        else:
            re_qc_top3.append({'model': m, 'defects': []})

    # 5. Top 3 defects per model - Re-Inspection MA
    re_ma_top3 = []
    for m in models_alphabetical:
        m_df = re_ma[re_ma[model_col] == m]
        sums = m_df[re_defect_cols].apply(pd.to_numeric, errors='coerce').sum()
        sums = sums[sums > 0]
        if not sums.empty:
            sums_df = pd.DataFrame({'qty': sums, 'site_tot': all_defect_sums.reindex(sums.index, fill_value=0)})
            sums_sorted = sums_df.sort_values(by=['qty', 'site_tot'], ascending=[False, False])
            re_ma_top3.append({'model': m, 'defects': list(sums_sorted.head(3)['qty'].items())})
        else:
            re_ma_top3.append({'model': m, 'defects': []})
        
    # 6. Top 3 defects per model - FTT (joined on Model / Model 2)
    ftt_top3 = []
    for m in models_alphabetical:
        m_df = ftt_df[ftt_df['Model'] == m]
        if len(m_df) == 0 and 'Model 2' in ftt_df.columns:
            m_df = ftt_df[ftt_df['Model 2'] == m]
        if len(m_df) > 0:
            sums = m_df.groupby('Defect Issues')["Issues Q'ty"].sum()
            sums = sums[sums > 0].sort_values(ascending=False)
            ftt_top3.append({'model': m, 'defects': list(sums.head(3).items())})
        else:
            ftt_top3.append({'model': m, 'defects': []})
        
    # 7. Top 3 defects per model - HFPA (Fail only, joined on Model 2 / Model)
    hfpa_defect_cols = defect_headers_after(hfpa_df, "Defect Q'ty (Total Defect Q'ty)")
    hfpa_fail = hfpa_df[hfpa_df['Pass/Fail'] == 'Fail']
    hfpa_top3 = []
    for m in models_alphabetical:
        m_df = hfpa_fail[hfpa_fail['Model 2'] == m]
        if len(m_df) == 0 and 'Model' in hfpa_fail.columns:
            m_df = hfpa_fail[hfpa_fail['Model'] == m]
        sums = m_df[hfpa_defect_cols].apply(pd.to_numeric, errors='coerce').sum()
        sums = sums[sums > 0].sort_values(ascending=False)
        hfpa_top3.append({'model': m, 'defects': list(sums.head(3).items())})
        
    return {
        'site': site,
        'site_name_full': SITE_NAMES_FULL.get(site, site),
        'total_records': total_records,
        'qc_total_fail': qc_total_fail,
        'ma_total_fail': ma_total_fail,
        'top_countries': top_countries,
        'country_counts_full': country_counts,
        'top_models': top5_models_tuples,
        'top5_models_list': top5_models,
        'models_alphabetical': models_alphabetical,
        'top5_defects': top5_defects,
        'site_all_defects': site_all_defects,
        'grand_total_5models': grand_total_5models,
        're_qc_top3': re_qc_top3,
        're_ma_top3': re_ma_top3,
        'ftt_top3': ftt_top3,
        'hfpa_top3': hfpa_top3
    }

def load_all_sites_data(re_dir, ftt_dir, hfpa_dir, log_fn=print):
    """Load and process data for all available sites."""
    all_sites_data = {}
    for site in SITES:
        re_f = find_site_file(re_dir, site)
        ftt_f = find_site_file(ftt_dir, site)
        hfpa_f = find_site_file(hfpa_dir, site)
        
        if not re_f or not ftt_f or not hfpa_f:
            log_fn(f"  [WARNING] Missing files for site {site}: Re-Ins={bool(re_f)}, FTT={bool(ftt_f)}, HFPA={bool(hfpa_f)}")
            continue
            
        log_fn(f"  Processing {site} ({SITE_NAMES_FULL.get(site, site)})...")
        all_sites_data[site] = extract_site_metrics(re_f, ftt_f, hfpa_f, site)
        log_fn(f"  ✓ {site}: QC Fails={all_sites_data[site]['qc_total_fail']}, MA={all_sites_data[site]['ma_total_fail']}")
        
    return all_sites_data


# ==============================================================================
#  EXCEL BUILDERS: RE-INSPECTION DATABASE
# ==============================================================================

def build_site_sheet_database(ws, data, site, color_map=None):
    """Build highly professional, tidy factory sheet (re-vh, re-vh2, RE-JV, RE-JV2)."""
    ws.views.sheetView[0].showGridLines = True
    
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 32
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 4
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 16
    ws.column_dimensions['H'].width = 16
    ws.column_dimensions['I'].width = 16
    
    ws.merge_cells('A2:I2')
    title_cell = ws.cell(2, 1, f"FACTORY AUDIT & RE-INSPECTION DASHBOARD — {site} ({SITE_NAMES_FULL.get(site, site)})")
    title_cell.font = font_title; title_cell.fill = fill_title; title_cell.alignment = align_center
    ws.row_dimensions[2].height = 28
    
    ws.cell(4, 1, "TOP DEFECTS ACROSS TOP 5 MODELS (QC + MA)")
    ws.merge_cells('A4:D4')
    ws.cell(4, 1).font = font_section; ws.cell(4, 1).fill = fill_section; ws.cell(4, 1).alignment = align_left
    
    headers_left = ['Defect Name', 'Original Metric', 'Defect Qty', 'Defect Share']
    for idx, h in enumerate(headers_left, 1):
        c = ws.cell(5, idx, h)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
    ws.row_dimensions[5].height = 22
    
    top5_def = data['top5_defects']
    gt = data['grand_total_5models']
    for idx, (defect_name, qty) in enumerate(top5_def):
        r = 6 + idx
        ws.row_dimensions[r].height = 20
        c1 = ws.cell(r, 1, f'=SUBSTITUTE(B{r},"Sum of ","")')
        c2 = ws.cell(r, 2, f'Sum of {defect_name}')
        c3 = ws.cell(r, 3, int(qty))
        c4 = ws.cell(r, 4, qty / gt if gt > 0 else 0)
        
        c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell
        c2.font = font_regular; c2.alignment = align_left; c2.border = border_cell
        c3.font = font_regular; c3.alignment = align_right; c3.border = border_cell; c3.number_format = '#,##0'
        c4.font = font_regular; c4.alignment = align_right; c4.border = border_cell; c4.number_format = '0.0%'
        
        has_custom_c1_fill = False
        if color_map:
            clean_d = str(defect_name).strip().replace('\xa0', ' ')
            d_hex = pptx_generator.get_defect_hex(clean_d, color_map)
            if d_hex and d_hex != 'DCDCDC':
                argb = 'FF' + d_hex
                c1.fill = PatternFill(fill_type='solid', start_color=argb, end_color=argb)
                rgb = pptx_generator.hex_to_rgb(d_hex)
                lum = 0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]
                c1.font = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF' if lum < 140 else '000000')
                has_custom_c1_fill = True

        if idx % 2 == 1:
            col_start = 2 if has_custom_c1_fill else 1
            for col_i in range(col_start, 5):
                ws.cell(r, col_i).fill = fill_zebra
                
    r_tot = 6 + len(top5_def)
    ws.cell(r_tot, 1, "Top 5 Defects Total")
    ws.cell(r_tot, 2, "")
    ws.cell(r_tot, 3, f"=SUM(C6:C{r_tot-1})")
    ws.cell(r_tot, 4, f"=SUM(D6:D{r_tot-1})")
    for col_i in range(1, 5):
        c = ws.cell(r_tot, col_i)
        c.font = font_bold; c.fill = fill_total; c.border = border_total
    ws.cell(r_tot, 1).alignment = align_left
    ws.cell(r_tot, 3).alignment = align_right; ws.cell(r_tot, 3).number_format = '#,##0'
    ws.cell(r_tot, 4).alignment = align_right; ws.cell(r_tot, 4).number_format = '0.0%'
    
    ws.cell(4, 6, "QC AUDIT FAIL BY DESTINATION COUNTRY")
    ws.merge_cells('F4:I4')
    ws.cell(4, 6).font = font_section; ws.cell(4, 6).fill = fill_section; ws.cell(4, 6).alignment = align_left
    
    headers_c = ['Country', 'Fail Qty', 'Total QC Fail', 'Fail Share']
    for idx, h in enumerate(headers_c, 6):
        c = ws.cell(5, idx, h)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
        
    qc_total = data['qc_total_fail']
    for idx, (country, count) in enumerate(data['top_countries']):
        r = 6 + idx
        ws.row_dimensions[r].height = 20
        c1 = ws.cell(r, 6, country)
        c2 = ws.cell(r, 7, count)
        c3 = ws.cell(r, 8, qc_total)
        c4 = ws.cell(r, 9, f'=G{r}/H{r}')
        
        c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell
        c2.font = font_regular; c2.alignment = align_right; c2.border = border_cell; c2.number_format = '#,##0'
        c3.font = font_regular; c3.alignment = align_right; c3.border = border_cell; c3.number_format = '#,##0'
        c4.font = font_regular; c4.alignment = align_right; c4.border = border_cell; c4.number_format = '0.0%'
        if idx % 2 == 1:
            for col_i in range(6, 10):
                ws.cell(r, col_i).fill = fill_zebra
                
    r_oth = 6 + len(data['top_countries'])
    c1 = ws.cell(r_oth, 6, "OTHERS")
    c2 = ws.cell(r_oth, 7, f'=H6-SUM(G6:G{r_oth-1})')
    c3 = ws.cell(r_oth, 8, '=H6')
    c4 = ws.cell(r_oth, 9, f'=G{r_oth}/H{r_oth}')
    for col_i in range(6, 10):
        c = ws.cell(r_oth, col_i)
        c.font = font_regular; c.border = border_cell
    c1.font = font_bold; c1.alignment = align_left
    c2.alignment = align_right; c2.number_format = '#,##0'
    c3.alignment = align_right; c3.number_format = '#,##0'
    c4.alignment = align_right; c4.number_format = '0.0%'
    
    r_m_sec = r_oth + 2
    ws.cell(r_m_sec, 6, "TOP 5 FAILED MODELS (QC AUDIT)")
    ws.merge_cells(f'F{r_m_sec}:I{r_m_sec}')
    ws.cell(r_m_sec, 6).font = font_section; ws.cell(r_m_sec, 6).fill = fill_section; ws.cell(r_m_sec, 6).alignment = align_left
    
    headers_m = ['Model Description', 'Fail Qty', 'Total QC Fail', 'Fail Share']
    r_m_head = r_m_sec + 1
    ws.row_dimensions[r_m_head].height = 22
    for idx, h in enumerate(headers_m, 6):
        c = ws.cell(r_m_head, idx, h)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
        
    for idx, (model, count) in enumerate(data['top_models']):
        r = r_m_head + 1 + idx
        ws.row_dimensions[r].height = 20
        c1 = ws.cell(r, 6, model)
        c2 = ws.cell(r, 7, count)
        c3 = ws.cell(r, 8, qc_total)
        c4 = ws.cell(r, 9, f'=G{r}/H{r}')
        
        c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell
        c2.font = font_regular; c2.alignment = align_right; c2.border = border_cell; c2.number_format = '#,##0'
        c3.font = font_regular; c3.alignment = align_right; c3.border = border_cell; c3.number_format = '#,##0'
        c4.font = font_regular; c4.alignment = align_right; c4.border = border_cell; c4.number_format = '0.0%'
        if idx % 2 == 1:
            for col_i in range(6, 10):
                ws.cell(r, col_i).fill = fill_zebra

def build_vn_indo_sheet(ws):
    """Build executive regional summary comparing Vietnam and Indonesia in English."""
    ws.views.sheetView[0].showGridLines = True
    
    for col_letter, w in [('B', 4), ('C', 22), ('D', 18), ('E', 18), ('F', 18), ('G', 6), ('H', 20), ('I', 18), ('J', 18)]:
        ws.column_dimensions[col_letter].width = w
        
    ws.merge_cells('C2:J2')
    title_cell = ws.cell(2, 3, "EXECUTIVE REGIONAL AUDIT SUMMARY — VIETNAM vs INDONESIA")
    title_cell.font = font_title; title_cell.fill = fill_title; title_cell.alignment = align_center
    ws.row_dimensions[2].height = 28
    
    ws.cell(4, 8, "REGIONAL COMPARISON (CLG)")
    ws.merge_cells('H4:J4')
    ws.cell(4, 8).font = font_section; ws.cell(4, 8).fill = fill_section; ws.cell(4, 8).alignment = align_left
    
    ws.cell(5, 8, "Region").font = font_tbl_header; ws.cell(5, 8).fill = fill_tbl_header; ws.cell(5, 8).alignment = align_center; ws.cell(5, 8).border = border_cell
    ws.cell(5, 9, "Total QC Fail").font = font_tbl_header; ws.cell(5, 9).fill = fill_tbl_header; ws.cell(5, 9).alignment = align_center; ws.cell(5, 9).border = border_cell
    ws.cell(5, 10, "Regional Share").font = font_tbl_header; ws.cell(5, 10).fill = fill_tbl_header; ws.cell(5, 10).alignment = align_center; ws.cell(5, 10).border = border_cell
    ws.row_dimensions[5].height = 22
    
    ws.cell(6, 8, "Vietnam (VN)").font = font_bold; ws.cell(6, 8).alignment = align_left; ws.cell(6, 8).border = border_cell
    ws.cell(6, 9, "=F6").font = font_regular; ws.cell(6, 9).alignment = align_right; ws.cell(6, 9).border = border_cell; ws.cell(6, 9).number_format = '#,##0'
    ws.cell(6, 10, "=I6/SUM(I6:I7)").font = font_regular; ws.cell(6, 10).alignment = align_right; ws.cell(6, 10).border = border_cell; ws.cell(6, 10).number_format = '0.0%'
    ws.row_dimensions[6].height = 20
    
    ws.cell(7, 8, "Indonesia (INDO)").font = font_bold; ws.cell(7, 8).alignment = align_left; ws.cell(7, 8).border = border_cell
    ws.cell(7, 9, "=F12").font = font_regular; ws.cell(7, 9).alignment = align_right; ws.cell(7, 9).border = border_cell; ws.cell(7, 9).number_format = '#,##0'
    ws.cell(7, 10, "=I7/SUM(I6:I7)").font = font_regular; ws.cell(7, 10).alignment = align_right; ws.cell(7, 10).border = border_cell; ws.cell(7, 10).number_format = '0.0%'
    ws.row_dimensions[7].height = 20
    
    ws.cell(8, 8, "TOTAL").font = font_bold; ws.cell(8, 8).alignment = align_left; ws.cell(8, 8).fill = fill_total; ws.cell(8, 8).border = border_total
    ws.cell(8, 9, "=SUM(I6:I7)").font = font_bold; ws.cell(8, 9).alignment = align_right; ws.cell(8, 9).fill = fill_total; ws.cell(8, 9).border = border_total; ws.cell(8, 9).number_format = '#,##0'
    ws.cell(8, 10, "100.0%").font = font_bold; ws.cell(8, 10).alignment = align_right; ws.cell(8, 10).fill = fill_total; ws.cell(8, 10).border = border_total
    ws.row_dimensions[8].height = 20
    
    ws.cell(4, 3, "VIETNAM FACTORIES BREAKDOWN")
    ws.merge_cells('C4:F4')
    ws.cell(4, 3).font = font_section; ws.cell(4, 3).fill = fill_section; ws.cell(4, 3).alignment = align_left
    
    headers_vn = ['Region', 'VH Factory', 'VH2 Factory', 'Total Vietnam']
    for idx, h in enumerate(headers_vn, 3):
        c = ws.cell(5, idx, h)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
        
    ws.cell(6, 3, "Fail Count").font = font_bold; ws.cell(6, 3).alignment = align_left; ws.cell(6, 3).border = border_cell
    ws.cell(6, 4, "='RE-VH'!H6").font = font_regular; ws.cell(6, 4).alignment = align_right; ws.cell(6, 4).border = border_cell; ws.cell(6, 4).number_format = '#,##0'
    ws.cell(6, 5, "='RE-VH2'!H6").font = font_regular; ws.cell(6, 5).alignment = align_right; ws.cell(6, 5).border = border_cell; ws.cell(6, 5).number_format = '#,##0'
    ws.cell(6, 6, "=D6+E6").font = font_bold; ws.cell(6, 6).alignment = align_right; ws.cell(6, 6).border = border_cell; ws.cell(6, 6).number_format = '#,##0'
    
    ws.cell(7, 3, "Percentage (%)").font = font_bold; ws.cell(7, 3).alignment = align_left; ws.cell(7, 3).fill = fill_total; ws.cell(7, 3).border = border_total
    ws.cell(7, 4, "=D6/F6").font = font_bold; ws.cell(7, 4).alignment = align_right; ws.cell(7, 4).fill = fill_total; ws.cell(7, 4).border = border_total; ws.cell(7, 4).number_format = '0.0%'
    ws.cell(7, 5, "=E6/F6").font = font_bold; ws.cell(7, 5).alignment = align_right; ws.cell(7, 5).fill = fill_total; ws.cell(7, 5).border = border_total; ws.cell(7, 5).number_format = '0.0%'
    ws.cell(7, 6, "100.0%").font = font_bold; ws.cell(7, 6).alignment = align_right; ws.cell(7, 6).fill = fill_total; ws.cell(7, 6).border = border_total
    
    ws.cell(10, 3, "INDONESIA FACTORIES BREAKDOWN")
    ws.merge_cells('C10:F10')
    ws.cell(10, 3).font = font_section; ws.cell(10, 3).fill = fill_section; ws.cell(10, 3).alignment = align_left
    
    headers_indo = ['Region', 'JV Factory', 'JV2 Factory', 'Total Indonesia']
    for idx, h in enumerate(headers_indo, 3):
        c = ws.cell(11, idx, h)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
    ws.row_dimensions[11].height = 22
    
    ws.cell(12, 3, "Fail Count").font = font_bold; ws.cell(12, 3).alignment = align_left; ws.cell(12, 3).border = border_cell
    ws.cell(12, 4, "='RE-JV'!H6").font = font_regular; ws.cell(12, 4).alignment = align_right; ws.cell(12, 4).border = border_cell; ws.cell(12, 4).number_format = '#,##0'
    ws.cell(12, 5, "='RE-JV2'!H6").font = font_regular; ws.cell(12, 5).alignment = align_right; ws.cell(12, 5).border = border_cell; ws.cell(12, 5).number_format = '#,##0'
    ws.cell(12, 6, "=D12+E12").font = font_bold; ws.cell(12, 6).alignment = align_right; ws.cell(12, 6).border = border_cell; ws.cell(12, 6).number_format = '#,##0'
    ws.row_dimensions[12].height = 20
    
    ws.cell(13, 3, "Percentage (%)").font = font_bold; ws.cell(13, 3).alignment = align_left; ws.cell(13, 3).fill = fill_total; ws.cell(13, 3).border = border_total
    ws.cell(13, 4, "=D12/F12").font = font_bold; ws.cell(13, 4).alignment = align_right; ws.cell(13, 4).fill = fill_total; ws.cell(13, 4).border = border_total; ws.cell(13, 4).number_format = '0.0%'
    ws.cell(13, 5, "=E12/F12").font = font_bold; ws.cell(13, 5).alignment = align_right; ws.cell(13, 5).fill = fill_total; ws.cell(13, 5).border = border_total; ws.cell(13, 5).number_format = '0.0%'
    ws.cell(13, 6, "100.0%").font = font_bold; ws.cell(13, 6).alignment = align_right; ws.cell(13, 6).fill = fill_total; ws.cell(13, 6).border = border_total
    ws.row_dimensions[13].height = 20
    
    ws.cell(15, 3, "* Note: QC Audit inspection failure numbers are dynamically linked from factory Re-Inspection worksheets.")
    ws.cell(15, 3).font = font_italic_note

def build_qc_or_ma_sheet(ws, all_sites_data, sheet_type='QC'):
    """Build tidy, professional QC or MA deep-dive sheet with styled master table and side cards."""
    ws.views.sheetView[0].showGridLines = True
    
    ws.column_dimensions['A'].width = 34.0
    ws.column_dimensions['B'].width = 14.0
    ws.column_dimensions['C'].width = 30.0
    ws.column_dimensions['D'].width = 14.0
    ws.column_dimensions['E'].width = 12.0
    ws.column_dimensions['F'].width = 4.0
    
    cols_grid = ['G', 'H', 'I', 'K', 'L', 'M', 'O', 'P', 'Q', 'S', 'T', 'U', 'W', 'X', 'Y']
    for c in cols_grid:
        ws.column_dimensions[c].width = 30.0 if c in ['G', 'K', 'O', 'S', 'W'] else 16.0
    for c in ['J', 'N', 'R', 'V']:
        ws.column_dimensions[c].width = 3.0
        
    ws.merge_cells('A2:E2')
    st_title = "QC AUDIT & FTT CROSS-ANALYSIS DASHBOARD" if sheet_type == 'QC' else "MA RE-AUDIT & HFPA CROSS-ANALYSIS DASHBOARD"
    t_cell = ws.cell(2, 1, st_title)
    t_cell.font = font_title; t_cell.fill = fill_title; t_cell.alignment = align_center
    ws.row_dimensions[2].height = 28
    
    ws.cell(4, 1, f"MASTER DEFECT AUDIT LOG — {sheet_type}")
    ws.merge_cells('A4:E4')
    ws.cell(4, 1).font = font_section; ws.cell(4, 1).fill = fill_section; ws.cell(4, 1).alignment = align_left
    
    headers_summary = ['Defect Description', 'Defect Qty', 'Shoe Model', 'Audit Station', 'Factory Site']
    for col_idx, h in enumerate(headers_summary, 1):
        c = ws.cell(5, col_idx, h)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
    ws.row_dimensions[5].height = 22
    
    summary_row = 6
    factory_order = ['VH', 'VH2', 'JV', 'JV2']
    
    grid_offsets = {
        'VH':  {'station1_start': 7,  'station2_start': 22, 'block_head1': 4,  'block_head2': 19},
        'VH2': {'station1_start': 37, 'station2_start': 52, 'block_head1': 34, 'block_head2': 49},
        'JV':  {'station1_start': 67, 'station2_start': 82, 'block_head1': 64, 'block_head2': 79},
        'JV2': {'station1_start': 97 if sheet_type == 'QC' else 101, 
                'station2_start': 112 if sheet_type == 'QC' else 116, 
                'block_head1': 94 if sheet_type == 'QC' else 98, 
                'block_head2': 109 if sheet_type == 'QC' else 113}
    }
    
    for factory in factory_order:
        if factory not in all_sites_data:
            continue
        data = all_sites_data[factory]
        offsets = grid_offsets[factory]
        
        st1_name = 'FTT' if sheet_type == 'QC' else 'HFPA'
        st2_name = 'Re-ins'
        
        st1_data = data['ftt_top3'] if sheet_type == 'QC' else data['hfpa_top3']
        st2_data = data['re_qc_top3'] if sheet_type == 'QC' else data['re_ma_top3']
        
        bh1 = offsets['block_head1']
        ws.cell(bh1, 7, f"FACTORY {factory} — {st1_name} TOP DEFECTS PER MODEL")
        ws.merge_cells(f"G{bh1}:Y{bh1}")
        ws.cell(bh1, 7).font = font_section; ws.cell(bh1, 7).fill = fill_section; ws.cell(bh1, 7).alignment = align_left
        ws.row_dimensions[bh1].height = 24
        
        subhead_row1 = offsets['station1_start'] - 1
        ws.row_dimensions[subhead_row1].height = 20
        for m_idx in range(5):
            base_col = 7 + m_idx * 4
            c1 = ws.cell(subhead_row1, base_col, 'Model Name')
            c2 = ws.cell(subhead_row1, base_col + 1, 'Defect Issue')
            c3 = ws.cell(subhead_row1, base_col + 2, 'Total Qty')
            for c in (c1, c2, c3):
                c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
                
        st1_start = offsets['station1_start']
        for m_idx, m_info in enumerate(st1_data):
            base_col = 7 + m_idx * 4
            model_name = m_info['model']
            defects = m_info['defects']
            
            for d_idx in range(3):
                curr_r = st1_start + d_idx
                ws.row_dimensions[curr_r].height = 19
                c1 = ws.cell(curr_r, base_col)
                c2 = ws.cell(curr_r, base_col + 1)
                c3 = ws.cell(curr_r, base_col + 2)
                
                c1.border = border_cell; c2.border = border_cell; c3.border = border_cell
                if d_idx < len(defects):
                    c1.value = model_name; c1.font = font_bold; c1.alignment = align_left
                    def_name, def_qty = defects[d_idx]
                    prefix = 'Sum of ' if sheet_type == 'MA' else ''
                    c2.value = f'{prefix}{def_name}'; c2.font = font_regular; c2.alignment = align_left
                    c3.value = int(def_qty); c3.font = font_regular; c3.alignment = align_right; c3.number_format = '#,##0'
                else:
                    c1.value = None; c2.value = None; c3.value = None
                    
        helper_start1 = st1_start + 3
        for m_idx in range(5):
            base_col = 7 + m_idx * 4
            col_letter_def = get_column_letter(base_col + 1)
            col_letter_qty = get_column_letter(base_col + 2)
            for d_idx in range(3):
                h_r = helper_start1 + d_idx
                src_r = st1_start + d_idx
                ws.cell(h_r, base_col + 1, f'=SUBSTITUTE({col_letter_def}{src_r},"Sum of ","")')
                ws.cell(h_r, base_col + 2, f'={col_letter_qty}{src_r}')
                
        st1_summary_start_row = summary_row
        for m_idx in range(5):
            base_col = 7 + m_idx * 4
            col_letter_model = get_column_letter(base_col)
            col_letter_def = get_column_letter(base_col + 1)
            col_letter_qty = get_column_letter(base_col + 2)
            
            for d_idx in range(3):
                h_r = helper_start1 + d_idx
                src_r = st1_start + d_idx
                ws.row_dimensions[summary_row].height = 20
                
                c1 = ws.cell(summary_row, 1, f'={col_letter_def}{h_r}')
                c2 = ws.cell(summary_row, 2, f'={col_letter_qty}{h_r}')
                c3 = ws.cell(summary_row, 3, f'={col_letter_model}{src_r}')
                c4 = ws.cell(summary_row, 4, st1_name)
                c5 = ws.cell(summary_row, 5, factory)
                
                c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell
                c2.font = font_regular; c2.alignment = align_right; c2.border = border_cell; c2.number_format = '#,##0'
                c3.font = font_regular; c3.alignment = align_left; c3.border = border_cell
                c4.font = font_bold; c4.alignment = align_center; c4.border = border_cell; c4.fill = fill_tbl_subheader
                c5.font = font_bold; c5.alignment = align_center; c5.border = border_cell
                
                if (summary_row - 6) % 2 == 1:
                    c1.fill = fill_zebra; c2.fill = fill_zebra; c3.fill = fill_zebra; c5.fill = fill_zebra
                summary_row += 1
                
        bh2 = offsets['block_head2']
        sec_title = f"FACTORY {factory} — RE-INSPECTION ({'INITIAL AUDIT - QC' if sheet_type == 'QC' else 'RE-AUDIT - MA'})"
        ws.cell(bh2, 7, sec_title)
        ws.merge_cells(f"G{bh2}:Y{bh2}")
        ws.cell(bh2, 7).font = font_section; ws.cell(bh2, 7).fill = fill_section; ws.cell(bh2, 7).alignment = align_left
        ws.row_dimensions[bh2].height = 24
        
        subhead_row2 = offsets['station2_start'] - 1
        ws.row_dimensions[subhead_row2].height = 20
        for m_idx in range(5):
            base_col = 7 + m_idx * 4
            c1 = ws.cell(subhead_row2, base_col, 'Model Name')
            c2 = ws.cell(subhead_row2, base_col + 1, 'Defect Issue')
            c3 = ws.cell(subhead_row2, base_col + 2, 'Defect Qty')
            for c in (c1, c2, c3):
                c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell
                
        st2_start = offsets['station2_start']
        for m_idx, m_info in enumerate(st2_data):
            base_col = 7 + m_idx * 4
            model_name = m_info['model']
            defects = m_info['defects']
            
            for d_idx in range(3):
                curr_r = st2_start + d_idx
                ws.row_dimensions[curr_r].height = 19
                c1 = ws.cell(curr_r, base_col)
                c2 = ws.cell(curr_r, base_col + 1)
                c3 = ws.cell(curr_r, base_col + 2)
                c1.border = border_cell; c2.border = border_cell; c3.border = border_cell
                
                if d_idx < len(defects):
                    c1.value = model_name; c1.font = font_bold; c1.alignment = align_left
                    def_name, def_qty = defects[d_idx]
                    c2.value = f'Sum of {def_name}'; c2.font = font_regular; c2.alignment = align_left
                    c3.value = int(def_qty); c3.font = font_regular; c3.alignment = align_right; c3.number_format = '#,##0'
                else:
                    c1.value = None; c2.value = None; c3.value = None
                    
        helper_start2 = st2_start + 3
        for m_idx in range(5):
            base_col = 7 + m_idx * 4
            col_letter_def = get_column_letter(base_col + 1)
            col_letter_qty = get_column_letter(base_col + 2)
            for d_idx in range(3):
                h_r = helper_start2 + d_idx
                src_r = st2_start + d_idx
                ws.cell(h_r, base_col + 1, f'=SUBSTITUTE({col_letter_def}{src_r},"Sum of ","")')
                ws.cell(h_r, base_col + 2, f'={col_letter_qty}{src_r}')
                
        for m_idx in range(5):
            base_col = 7 + m_idx * 4
            col_letter_def = get_column_letter(base_col + 1)
            col_letter_qty = get_column_letter(base_col + 2)
            
            for d_idx in range(3):
                h_r = helper_start2 + d_idx
                src_r = st2_start + d_idx
                relative_idx = (m_idx * 3) + d_idx
                st1_corresp_row = st1_summary_start_row + relative_idx
                ws.row_dimensions[summary_row].height = 20
                
                c1 = ws.cell(summary_row, 1, f'={col_letter_def}{h_r}')
                c2 = ws.cell(summary_row, 2, f'={col_letter_qty}{h_r}')
                c3 = ws.cell(summary_row, 3, f'=C{st1_corresp_row}')
                c4 = ws.cell(summary_row, 4, st2_name)
                c5 = ws.cell(summary_row, 5, factory)
                
                c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell
                c2.font = font_regular; c2.alignment = align_right; c2.border = border_cell; c2.number_format = '#,##0'
                c3.font = font_regular; c3.alignment = align_left; c3.border = border_cell
                c4.font = font_bold; c4.alignment = align_center; c4.border = border_cell; c4.fill = fill_total
                c5.font = font_bold; c5.alignment = align_center; c5.border = border_cell
                
                if (summary_row - 6) % 2 == 1:
                    c1.fill = fill_zebra; c2.fill = fill_zebra; c3.fill = fill_zebra; c5.fill = fill_zebra
                summary_row += 1


# ==============================================================================
#  EXCEL BUILDERS: RECYCLE REPORT
# ==============================================================================

def _style_header_row_recycle(ws, row, col_start, headers):
    for idx, h in enumerate(headers):
        c = ws.cell(row, col_start + idx, h)
        c.font = font_header_recycle; c.fill = fill_tbl_header; c.alignment = align_center; c.border = border_cell

def _style_section_band_recycle(ws, row, col, text, merge_end_col=None):
    cell = ws.cell(row, col, text)
    cell.font = font_section_recycle; cell.fill = fill_section_recycle; cell.alignment = align_left
    if merge_end_col and merge_end_col > col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_end_col)

def write_data_row_recycle(ws, row, col, values, fonts=None, aligns=None, num_fmts=None):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.border = border_cell
        c.font = fonts[i] if (fonts and i < len(fonts) and fonts[i]) else font_regular
        c.alignment = aligns[i] if (aligns and i < len(aligns) and aligns[i]) else align_left
        if num_fmts and i < len(num_fmts) and num_fmts[i]:
            c.number_format = num_fmts[i]

def build_site_sheet_recycle(ws, data, site, report_title):
    """Build a single site sheet matching the Recycle_Report format."""
    ws.sheet_view.showGridLines = True

    ws.column_dimensions['A'].width = 42
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 14

    row = 1
    title_cell = ws.cell(row, 1, report_title)
    title_cell.font = font_title_recycle; title_cell.alignment = align_left
    row += 2

    # SECTION 1: Top 5 countries by fail qty (QC)
    _style_section_band_recycle(ws, row, 1, "Top 5 countries by fail qty (Re-Inspection, QC)", 2)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Country', 'Fail Qty'])
    row += 1
    for country, qty in data['top_countries']:
        write_data_row_recycle(ws, row, 1, [country, int(qty)], aligns=[align_left, align_right], num_fmts=[None, '#,##0'])
        row += 1
    row += 2

    # SECTION 2: Top 5 models by fail qty (QC)
    _style_section_band_recycle(ws, row, 1, "Top 5 models by fail qty (Re-Inspection, QC)", 2)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Model', 'Fail Qty'])
    row += 1
    for model, qty in data['top_models']:
        write_data_row_recycle(ws, row, 1, [model, int(qty)], aligns=[align_left, align_right], num_fmts=[None, '#,##0'])
        row += 1
    row += 2

    # SECTION 3: Top 5 defects across top 5 models (QC+MA)
    _style_section_band_recycle(ws, row, 1, "Top 5 defects across top 5 models (Re-Inspection, QC+MA)", 4)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Defect', 'Defect Qty', 'Total Defect Qty (5 models)', 'TL%'])
    row += 1
    gt = data['grand_total_5models']
    for defect, qty in data['top5_defects']:
        pct = qty / gt if gt > 0 else 0
        write_data_row_recycle(ws, row, 1, [defect, int(qty), int(gt), pct], aligns=[align_left, align_right, align_right, align_right], num_fmts=[None, '#,##0', '#,##0', '0.0%'])
        row += 1
    row += 2

    # SECTION 4: Top 3 defects per model — Re-Inspection QC
    _style_section_band_recycle(ws, row, 1, "Top 3 defects per model - Re-Inspection QC (initial audit)", 3)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Model', 'Defect', 'Qty'])
    row += 1
    for m_info in data['re_qc_top3']:
        for defect, qty in m_info['defects']:
            write_data_row_recycle(ws, row, 1, [m_info['model'], defect, int(qty)], aligns=[align_left, align_left, align_right], num_fmts=[None, None, '#,##0'])
            row += 1
    row += 2

    # SECTION 5: Top 3 defects per model — Re-Inspection MA
    _style_section_band_recycle(ws, row, 1, "Top 3 defects per model - Re-Inspection MA (re-audit)", 3)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Model', 'Defect', 'Qty'])
    row += 1
    for m_info in data['re_ma_top3']:
        for defect, qty in m_info['defects']:
            write_data_row_recycle(ws, row, 1, [m_info['model'], defect, int(qty)], aligns=[align_left, align_left, align_right], num_fmts=[None, None, '#,##0'])
            row += 1
    row += 2

    # SECTION 6: Top 3 defects per model — FTT
    _style_section_band_recycle(ws, row, 1, "Top 3 defects per model - FTT (Model field, full month volume)", 3)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Model', 'Defect', 'Qty'])
    row += 1
    for m_info in data['ftt_top3']:
        for defect, qty in m_info['defects']:
            write_data_row_recycle(ws, row, 1, [m_info['model'], defect, int(qty)], aligns=[align_left, align_left, align_right], num_fmts=[None, None, '#,##0'])
            row += 1
    row += 2

    # SECTION 7: Top 3 defects per model — HFPA (Fail only)
    _style_section_band_recycle(ws, row, 1, "Top 3 defects per model - HFPA (Fail only)", 3)
    row += 1
    _style_header_row_recycle(ws, row, 1, ['Model', 'Defect', 'Qty'])
    row += 1
    for m_info in data['hfpa_top3']:
        for defect, qty in m_info['defects']:
            write_data_row_recycle(ws, row, 1, [m_info['model'], defect, int(qty)], aligns=[align_left, align_left, align_right], num_fmts=[None, None, '#,##0'])
            row += 1


# ==============================================================================
#  HIGH-LEVEL GENERATION PIPELINES
# ==============================================================================

def generate_database_workbook(re_dir, ftt_dir, hfpa_dir, output_file=None, month_label=None, color_template="Color_Template.xlsx", log_fn=print):
    """Generate the full Re-Inspection-Database.xlsx workbook."""
    if month_label is None:
        month_label = detect_reporting_period(re_dir)
    if output_file is None or output_file == "Re-Inspection-Database.xlsx":
        output_file = format_db_filename(month_label)
        
    log_fn(f"Starting Re-Inspection Master Database export ({month_label}) to: {output_file}")
    all_sites_data = load_all_sites_data(re_dir, ftt_dir, hfpa_dir, log_fn)
    if not all_sites_data:
        raise ValueError("No factory site data could be loaded. Check directory paths.")

    color_map = {}
    if color_template and os.path.exists(color_template):
        all_defects = set()
        for s_data in all_sites_data.values():
            for d in s_data.get('top5_defects', []):
                if d and d[0]:
                    all_defects.add(str(d[0]).strip().replace('\xa0', ' '))
            for grp in ['ftt_top3', 'hfpa_top3', 're_qc_top3', 're_ma_top3']:
                for m_item in s_data.get(grp, []):
                    for d_item in m_item.get('defects', []):
                        if d_item and d_item[0]:
                            all_defects.add(str(d_item[0]).strip().replace('\xa0', ' '))
        color_map = pptx_generator.sync_and_register_defects(all_defects, color_template, log_fn=log_fn)
        
    wb = openpyxl.Workbook()
    
    ws_revh = wb.active
    ws_revh.title = 'RE-VH'
    if 'VH' in all_sites_data:
        build_site_sheet_database(ws_revh, all_sites_data['VH'], 'VH', color_map=color_map)
    
    if 'VH2' in all_sites_data:
        ws_revh2 = wb.create_sheet('RE-VH2')
        build_site_sheet_database(ws_revh2, all_sites_data['VH2'], 'VH2', color_map=color_map)
        
    if 'JV' in all_sites_data:
        ws_rejv = wb.create_sheet('RE-JV')
        build_site_sheet_database(ws_rejv, all_sites_data['JV'], 'JV', color_map=color_map)
        
    if 'JV2' in all_sites_data:
        ws_rejv2 = wb.create_sheet('RE-JV2')
        build_site_sheet_database(ws_rejv2, all_sites_data['JV2'], 'JV2', color_map=color_map)
        
    ws_vnindo = wb.create_sheet('VN-INDO')
    build_vn_indo_sheet(ws_vnindo)
    
    ws_qc = wb.create_sheet('QC')
    build_qc_or_ma_sheet(ws_qc, all_sites_data, 'QC')
    
    ws_ma = wb.create_sheet('MA')
    build_qc_or_ma_sheet(ws_ma, all_sites_data, 'MA')
    
    out_dir = os.path.dirname(os.path.abspath(output_file))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    try:
        wb.save(output_file)
        log_fn(f"[OK] Master Database generated successfully: {os.path.abspath(output_file)}")
        return output_file
    except PermissionError:
        base, ext = os.path.splitext(output_file)
        alt_path = f"{base}_updated{ext}"
        wb.save(alt_path)
        log_fn(f"[WARNING] {output_file} is locked by Excel/WPS. Saved to {alt_path} instead.")
        return alt_path

def generate_recycle_workbook(re_dir, ftt_dir, hfpa_dir, output_file=None, month_label=None, log_fn=print):
    """Generate the Recycle_Report.xlsx workbook."""
    if month_label is None:
        month_label = detect_reporting_period(re_dir)
    if output_file is None or output_file == "Recycle_Report.xlsx":
        output_file = format_recycle_filename(month_label)
        
    log_fn(f"Starting Recycle Report export ({month_label}) to: {output_file}")
    all_sites_data = load_all_sites_data(re_dir, ftt_dir, hfpa_dir, log_fn)
    if not all_sites_data:
        raise ValueError("No factory site data could be loaded. Check directory paths.")
        
    wb = openpyxl.Workbook()
    first = True
    for site in SITES:
        if site not in all_sites_data:
            continue
        if first:
            ws = wb.active
            ws.title = site
            first = False
        else:
            ws = wb.create_sheet(site)
            
        title = f"Recycle report - {site} - {month_label}"
        build_site_sheet_recycle(ws, all_sites_data[site], site, title)
        
    out_dir = os.path.dirname(os.path.abspath(output_file))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_file)
    log_fn(f"✓ Recycle Report generated successfully: {os.path.abspath(output_file)}")
    return output_file

def generate_presentation_report(re_dir, ftt_dir, hfpa_dir, template_pptx="RE-INS REPORT.APR.2026.pptx", output_pptx=None, month_label=None, color_template="Color_Template.xlsx", log_fn=print):
    """Generate the executive PowerPoint report with exact layout and colors."""
    if month_label is None:
        month_label = detect_reporting_period(re_dir)
    if output_pptx is None or output_pptx == "Re-Ins Report-June-2026.pptx":
        output_pptx = format_pptx_filename(month_label)
    all_sites_data = load_all_sites_data(re_dir, ftt_dir, hfpa_dir, log_fn)
    return pptx_generator.export_reinspection_presentation(all_sites_data, template_pptx, output_pptx, month_label, color_template_path=color_template, log_fn=log_fn)

def generate_all_reports(re_dir, ftt_dir, hfpa_dir, db_output=None, recycle_output=None, pptx_output=None, template_pptx="RE-INS REPORT.APR.2026.pptx", color_template="Color_Template.xlsx", log_fn=print):
    """Generate all workbooks and PowerPoint presentation."""
    period = detect_reporting_period(re_dir)
    if db_output is None or db_output == "Re-Inspection-Database.xlsx":
        db_output = format_db_filename(period)
    if recycle_output is None or recycle_output == "Recycle_Report.xlsx":
        recycle_output = format_recycle_filename(period)
    if pptx_output is None:
        pptx_output = format_pptx_filename(period)
    generate_database_workbook(re_dir, ftt_dir, hfpa_dir, db_output, month_label=period, color_template=color_template, log_fn=log_fn)
    generate_recycle_workbook(re_dir, ftt_dir, hfpa_dir, recycle_output, month_label=period, log_fn=log_fn)
    if os.path.exists(template_pptx):
        generate_presentation_report(re_dir, ftt_dir, hfpa_dir, template_pptx, pptx_output, month_label=period, color_template=color_template, log_fn=log_fn)
    log_fn(f"[OK] ALL reports generated successfully!\n  - Database: {db_output}\n  - Recycle: {recycle_output}\n  - Presentation: {pptx_output}")


# ==============================================================================
#  NATIVE PYTHON GUI DASHBOARD (TKINTER + CUSTOM CANVAS CHARTS)
# ==============================================================================

def launch_dashboard(default_re="Re-Inspection", default_ftt="FTT", default_hfpa="HFPA", default_template="RE-INS REPORT.APR.2026.pptx", default_color_tpl="Color_Template.xlsx"):
    """Launch the sleek native Python GUI visualizer dashboard."""
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    root = tk.Tk()
    root.title("Re-Inspection & Quality Intelligence Suite")
    root.geometry("1240x840")
    root.minsize(1050, 720)

    style = ttk.Style()
    style.theme_use('clam')
    
    style.configure('TNotebook', background='#EBF0F5', tabmargins=[5, 5, 0, 0])
    style.configure('TNotebook.Tab', background='#D0DCE5', foreground='#1F4E78', font=('Segoe UI', 10, 'bold'), padding=[16, 6])
    style.map('TNotebook.Tab', background=[('selected', '#FFFFFF')], foreground=[('selected', '#1F4E78')])
    style.configure('Treeview', font=('Segoe UI', 9), rowheight=24)
    style.configure('Treeview.Heading', font=('Segoe UI', 9, 'bold'), background='#E2E8F0', foreground='#1F4E78')

    cur_dir = os.getcwd()
    out_dir = os.path.join(cur_dir, DEFAULT_OUTPUT_DIR)
    os.makedirs(out_dir, exist_ok=True)
    init_period = detect_reporting_period(default_re)
    re_var = tk.StringVar(value=os.path.join(cur_dir, default_re) if os.path.exists(os.path.join(cur_dir, default_re)) else default_re)
    ftt_var = tk.StringVar(value=os.path.join(cur_dir, default_ftt) if os.path.exists(os.path.join(cur_dir, default_ftt)) else default_ftt)
    hfpa_var = tk.StringVar(value=os.path.join(cur_dir, default_hfpa) if os.path.exists(os.path.join(cur_dir, default_hfpa)) else default_hfpa)
    tpl_var = tk.StringVar(value=os.path.join(cur_dir, default_template) if os.path.exists(os.path.join(cur_dir, default_template)) else default_template)
    color_tpl_var = tk.StringVar(value=os.path.join(cur_dir, default_color_tpl) if os.path.exists(os.path.join(cur_dir, default_color_tpl)) else default_color_tpl)
    
    db_out_var = tk.StringVar(value=os.path.join(cur_dir, format_db_filename(init_period)))
    recycle_out_var = tk.StringVar(value=os.path.join(cur_dir, format_recycle_filename(init_period)))
    pptx_out_var = tk.StringVar(value=os.path.join(cur_dir, format_pptx_filename(init_period)))
    status_msg_var = tk.StringVar(value="Ready to load data")
    period_var = tk.StringVar(value=f"Period: {init_period}")
    
    loaded_data = {'sites': {}, 'period': ''}

    # Top Executive Banner
    header = tk.Frame(root, bg='#1F4E78', height=75)
    header.pack(fill=tk.X, side=tk.TOP)
    header.pack_propagate(False)

    title_frame = tk.Frame(header, bg='#1F4E78')
    title_frame.pack(side=tk.LEFT, padx=20, pady=10)
    
    tk.Label(title_frame, text="RE-INSPECTION & QUALITY INTELLIGENCE SUITE", font=('Segoe UI', 15, 'bold'), fg='#FFFFFF', bg='#1F4E78').pack(anchor='w')
    tk.Label(title_frame, text="Executive Visualizer & Automated Excel/PPTX Report Generator (FTT • HFPA • Re-Ins)", font=('Segoe UI', 9), fg='#D9E1F2', bg='#1F4E78').pack(anchor='w')

    badge_frame = tk.Frame(header, bg='#1F4E78')
    badge_frame.pack(side=tk.RIGHT, padx=20, pady=10)
    
    lbl_period = tk.Label(badge_frame, textvariable=period_var, font=('Segoe UI', 10, 'bold'), fg='#1F4E78', bg='#D9E1F2', padx=12, pady=4, relief=tk.FLAT)
    lbl_period.pack(anchor='e')
    
    lbl_status_badge = tk.Label(badge_frame, textvariable=status_msg_var, font=('Segoe UI', 8), fg='#A6C8FF', bg='#1F4E78')
    lbl_status_badge.pack(anchor='e', pady=2)

    # Control Ribbon
    ribbon = tk.Frame(root, bg='#F0F4F8', bd=1, relief=tk.SOLID, padx=12, pady=8)
    ribbon.pack(fill=tk.X, padx=10, pady=(8, 4))

    r1 = tk.Frame(ribbon, bg='#F0F4F8')
    r1.pack(fill=tk.X)

    def browse_dir(v, title):
        p = filedialog.askdirectory(title=title)
        if p: v.set(p)

    def browse_file(v, title, ext, filetypes):
        p = filedialog.askopenfilename(title=title, defaultextension=ext, filetypes=filetypes)
        if p: v.set(p)

    def make_ribbon_picker(parent, label, var, title):
        f = tk.Frame(parent, bg='#F0F4F8')
        f.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=4)
        tk.Label(f, text=label, font=('Segoe UI', 8, 'bold'), fg='#2F5597', bg='#F0F4F8').pack(anchor='w')
        box = tk.Frame(f, bg='#F0F4F8')
        box.pack(fill=tk.X)
        e = tk.Entry(box, textvariable=var, font=('Segoe UI', 9), bg='#FFFFFF', relief=tk.SOLID, bd=1)
        e.pack(side=tk.LEFT, fill=tk.X, expand=True)
        b = tk.Button(box, text="...", command=lambda: browse_dir(var, title), width=3, bg='#E2E8F0', font=('Segoe UI', 8), relief=tk.GROOVE)
        b.pack(side=tk.RIGHT, padx=(2, 0))

    make_ribbon_picker(r1, "1. Re-Inspection Dir", re_var, "Select Re-Inspection Folder")
    make_ribbon_picker(r1, "2. FTT Dir", ftt_var, "Select FTT Folder")
    make_ribbon_picker(r1, "3. HFPA Dir", hfpa_var, "Select HFPA Folder")

    # Action Buttons on Ribbon
    btn_box = tk.Frame(ribbon, bg='#F0F4F8')
    btn_box.pack(fill=tk.X, pady=(8, 0))

    btn_reload = tk.Button(btn_box, text="🔄 Reload & Visualize", bg='#2F5597', fg='white', font=('Segoe UI', 9, 'bold'), padx=10, pady=4, relief=tk.RAISED)
    btn_reload.pack(side=tk.LEFT, padx=3)

    btn_gen_db = tk.Button(btn_box, text="📘 Export Database", bg='#1F4E78', fg='white', font=('Segoe UI', 9, 'bold'), padx=8, pady=4, relief=tk.RAISED)
    btn_gen_db.pack(side=tk.LEFT, padx=3)

    btn_gen_recycle = tk.Button(btn_box, text="📗 Export Recycle Report", bg='#107C41', fg='white', font=('Segoe UI', 9, 'bold'), padx=8, pady=4, relief=tk.RAISED)
    btn_gen_recycle.pack(side=tk.LEFT, padx=3)

    btn_gen_pptx = tk.Button(btn_box, text="📊 Export PPTX Presentation", bg='#D9531E', fg='white', font=('Segoe UI', 9, 'bold'), padx=10, pady=4, relief=tk.RAISED)
    btn_gen_pptx.pack(side=tk.LEFT, padx=3)

    btn_gen_both = tk.Button(btn_box, text="🚀 Export All", bg='#0F2840', fg='white', font=('Segoe UI', 9, 'bold'), padx=10, pady=4, relief=tk.RAISED)
    btn_gen_both.pack(side=tk.LEFT, padx=3)

    btn_open_folder = tk.Button(btn_box, text="📂 Open Output Folder", bg='#E2E8F0', fg='#1F4E78', font=('Segoe UI', 9, 'bold'), padx=8, pady=4, relief=tk.GROOVE, command=lambda: os.startfile(os.path.join(cur_dir, DEFAULT_OUTPUT_DIR) if os.path.exists(os.path.join(cur_dir, DEFAULT_OUTPUT_DIR)) else cur_dir))
    btn_open_folder.pack(side=tk.RIGHT, padx=4)

    # Main Notebook (Tabs)
    notebook = ttk.Notebook(root)
    notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

    tab_overview = tk.Frame(notebook, bg='#F4F7FB')
    tab_factory = tk.Frame(notebook, bg='#F4F7FB')
    tab_regional = tk.Frame(notebook, bg='#F4F7FB')
    tab_cross = tk.Frame(notebook, bg='#F4F7FB')
    tab_logs = tk.Frame(notebook, bg='#F4F7FB')

    notebook.add(tab_overview, text=" 📊 Executive Overview ")
    notebook.add(tab_factory, text=" 🏭 Factory Deep-Dive ")
    notebook.add(tab_regional, text=" 🌏 VN vs INDO Summary ")
    notebook.add(tab_cross, text=" 🔬 Station Cross-Analysis ")
    notebook.add(tab_logs, text=" 📑 Exporter & Console Logs ")

    # Custom Canvas Chart Drawing Utilities
    def draw_card(canvas, x, y, w, h, title, val, subtext, color='#1F4E78', bg='#FFFFFF'):
        canvas.create_rectangle(x, y, x + w, y + h, fill=bg, outline='#D1D5DB', width=1)
        canvas.create_rectangle(x, y, x + 6, y + h, fill=color, outline='')
        canvas.create_text(x + 16, y + 16, text=title.upper(), font=('Segoe UI', 8, 'bold'), fill='#64748B', anchor='w')
        canvas.create_text(x + 16, y + 42, text=str(val), font=('Segoe UI', 18, 'bold'), fill='#0F172A', anchor='w')
        if subtext:
            canvas.create_text(x + 16, y + 68, text=subtext, font=('Segoe UI', 8), fill='#64748B', anchor='w')

    def draw_bar_chart(canvas, x, y, w, h, title, categories, values, colors=None, format_pct=False):
        canvas.create_rectangle(x, y, x + w, y + h, fill='#FFFFFF', outline='#E2E8F0', width=1)
        canvas.create_text(x + 14, y + 16, text=title, font=('Segoe UI', 10, 'bold'), fill='#1F4E78', anchor='w')
        if not categories or not values or sum(values) == 0:
            canvas.create_text(x + w / 2, y + h / 2, text="No Data Available", font=('Segoe UI', 10, 'italic'), fill='#94A3B8')
            return

        total = sum(values)
        max_v = max(values)
        bar_count = len(categories)
        palette = colors or ['#1F4E78', '#2F5597', '#4172B8', '#6BA4E8', '#94C5FF', '#BBD7FC']
        
        top_pad = 38
        avail_h = h - top_pad - 15
        row_h = avail_h / max(bar_count, 1)
        max_bar_w = w - 240
        
        for i, (cat, val) in enumerate(zip(categories, values)):
            cy = y + top_pad + i * row_h + row_h / 2
            bar_w = (val / max_v) * max_bar_w if max_v > 0 else 0
            color = palette[i % len(palette)]
            
            label_text = str(cat)
            if len(label_text) > 24: label_text = label_text[:22] + "..."
            canvas.create_text(x + 14, cy, text=label_text, font=('Segoe UI', 8, 'bold'), fill='#334155', anchor='w')
            
            bx = x + 150
            by1 = cy - 8
            by2 = cy + 8
            canvas.create_rectangle(bx, by1, bx + max(bar_w, 2), by2, fill=color, outline='')
            
            pct_str = f"({(val/total*100):.1f}%)" if total > 0 else ""
            val_str = f"{int(val):,}" if not format_pct else f"{val:.1%}"
            canvas.create_text(bx + bar_w + 8, cy, text=f"{val_str} {pct_str}", font=('Segoe UI', 8), fill='#475569', anchor='w')

    def draw_donut_chart(canvas, cx, cy, radius, title, labels, values, colors=None):
        canvas.create_text(cx, cy - radius - 18, text=title, font=('Segoe UI', 10, 'bold'), fill='#1F4E78', anchor='center')
        if not labels or not values or sum(values) == 0:
            canvas.create_text(cx, cy, text="No Data", font=('Segoe UI', 10, 'italic'), fill='#94A3B8')
            return
            
        total = sum(values)
        start_angle = 90
        palette = colors or ['#1F4E78', '#2F5597', '#3B82F6', '#10B981', '#F59E0B', '#EF4444', '#8B5CF6']
        
        for i, (lbl, val) in enumerate(zip(labels, values)):
            extent = (val / total) * 360.0
            col = palette[i % len(palette)]
            canvas.create_arc(cx - radius, cy - radius, cx + radius, cy + radius,
                              start=start_angle, extent=-extent, fill=col, outline='#FFFFFF', width=2)
            start_angle -= extent
            
        hole_r = radius * 0.55
        canvas.create_oval(cx - hole_r, cy - hole_r, cx + hole_r, cy + hole_r, fill='#FFFFFF', outline='')
        canvas.create_text(cx, cy - 6, text="TOTAL", font=('Segoe UI', 7, 'bold'), fill='#64748B')
        canvas.create_text(cx, cy + 10, text=f"{int(total):,}", font=('Segoe UI', 11, 'bold'), fill='#0F172A')
        
        leg_y = cy + radius + 14
        for i, (lbl, val) in enumerate(zip(labels[:5], values[:5])):
            col = palette[i % len(palette)]
            pct = (val / total * 100) if total > 0 else 0
            lx = cx - radius
            ly = leg_y + i * 16
            canvas.create_rectangle(lx, ly, lx + 10, ly + 10, fill=col, outline='')
            canvas.create_text(lx + 16, ly + 5, text=f"{lbl}: {int(val):,} ({pct:.1f}%)", font=('Segoe UI', 8), fill='#334155', anchor='w')

    # TAB 1: EXECUTIVE OVERVIEW
    ov_scroll = tk.Frame(tab_overview, bg='#F4F7FB')
    ov_scroll.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    canvas_kpis = tk.Canvas(ov_scroll, bg='#F4F7FB', height=95, highlightthickness=0)
    canvas_kpis.pack(fill=tk.X, side=tk.TOP, pady=(0, 10))

    canvas_charts_top = tk.Canvas(ov_scroll, bg='#F4F7FB', height=280, highlightthickness=0)
    canvas_charts_top.pack(fill=tk.X, side=tk.TOP, pady=(0, 10))

    canvas_charts_bottom = tk.Canvas(ov_scroll, bg='#F4F7FB', height=280, highlightthickness=0)
    canvas_charts_bottom.pack(fill=tk.BOTH, expand=True)

    def render_overview(all_data):
        canvas_kpis.delete('all')
        canvas_charts_top.delete('all')
        canvas_charts_bottom.delete('all')

        if not all_data:
            canvas_kpis.create_text(300, 45, text="Please click 'Reload & Visualize' to populate metrics.", font=('Segoe UI', 11, 'italic'), fill='#64748B')
            return

        total_qc = sum(d['qc_total_fail'] for d in all_data.values())
        total_ma = sum(d['ma_total_fail'] for d in all_data.values())
        total_rec = sum(d['total_records'] for d in all_data.values())
        
        vn_qc = sum(all_data[s]['qc_total_fail'] for s in ['VH', 'VH2'] if s in all_data)
        indo_qc = sum(all_data[s]['qc_total_fail'] for s in ['JV', 'JV2'] if s in all_data)
        
        all_def_dict = {}
        for d in all_data.values():
            for def_name, qty in d['top5_defects']:
                all_def_dict[def_name] = all_def_dict.get(def_name, 0) + qty
        top_def_sorted = sorted(all_def_dict.items(), key=lambda x: x[1], reverse=True)
        top_def_name = top_def_sorted[0][0] if top_def_sorted else "N/A"
        top_def_qty = top_def_sorted[0][1] if top_def_sorted else 0

        w_kpi = 220
        gap = 16
        draw_card(canvas_kpis, 10 + 0 * (w_kpi + gap), 10, w_kpi, 75, "Total QC Failures", f"{total_qc:,}", f"From {total_rec:,} Audit Rows", '#1F4E78')
        draw_card(canvas_kpis, 10 + 1 * (w_kpi + gap), 10, w_kpi, 75, "MA Re-Audits", f"{total_ma:,}", "Secondary Re-inspection", '#2F5597')
        draw_card(canvas_kpis, 10 + 2 * (w_kpi + gap), 10, w_kpi, 75, "Top Defect Issue", str(int(top_def_qty)), f"{top_def_name[:18]}...", '#DC2626')
        draw_card(canvas_kpis, 10 + 3 * (w_kpi + gap), 10, w_kpi, 75, "Vietnam (VH+VH2)", f"{vn_qc:,}", f"{(vn_qc/total_qc*100):.1f}% of total" if total_qc else "0%", '#10B981')
        draw_card(canvas_kpis, 10 + 4 * (w_kpi + gap), 10, w_kpi, 75, "Indonesia (JV+JV2)", f"{indo_qc:,}", f"{(indo_qc/total_qc*100):.1f}% of total" if total_qc else "0%", '#F59E0B')

        site_names = [s for s in SITES if s in all_data]
        site_qc = [all_data[s]['qc_total_fail'] for s in site_names]
        draw_bar_chart(canvas_charts_top, 10, 10, 600, 260, "Factory QC Failure Comparison", [f"{s} ({SITE_NAMES_FULL.get(s,s)})" for s in site_names], site_qc)

        reg_labels = ['Vietnam (VH/VH2)', 'Indonesia (JV/JV2)']
        reg_vals = [vn_qc, indo_qc]
        draw_donut_chart(canvas_charts_top, 910, 110, 75, "Regional QC Fail Distribution", reg_labels, reg_vals, ['#1F4E78', '#F59E0B'])

        top_def_cats = [d[0] for d in top_def_sorted[:6]]
        top_def_vals = [d[1] for d in top_def_sorted[:6]]
        draw_bar_chart(canvas_charts_bottom, 10, 10, 600, 260, "Top Defect Issues Across All Sites (QC + MA)", top_def_cats, top_def_vals, ['#1F4E78', '#2563EB', '#3B82F6', '#60A5FA', '#93C5FD', '#BFDBFE'])

        all_country_dict = {}
        for d in all_data.values():
            for c, q in d['top_countries']:
                all_country_dict[c] = all_country_dict.get(c, 0) + q
        top_c_sorted = sorted(all_country_dict.items(), key=lambda x: x[1], reverse=True)[:5]
        draw_bar_chart(canvas_charts_bottom, 630, 10, 560, 260, "Top Destination Countries by QC Failures", [c[0] for c in top_c_sorted], [c[1] for c in top_c_sorted], ['#047857', '#059669', '#10B981', '#34D399', '#6EE7B7'])

    # TAB 2: FACTORY DEEP DIVE
    factory_top = tk.Frame(tab_factory, bg='#F4F7FB', pady=8)
    factory_top.pack(fill=tk.X, padx=12)

    tk.Label(factory_top, text="Select Factory Site:", font=('Segoe UI', 10, 'bold'), fg='#1F4E78', bg='#F4F7FB').pack(side=tk.LEFT, padx=(0, 8))
    site_combo = ttk.Combobox(factory_top, values=SITES, state='readonly', font=('Segoe UI', 9), width=10)
    site_combo.set('VH')
    site_combo.pack(side=tk.LEFT, padx=4)

    lbl_site_full = tk.Label(factory_top, text="Vietnam Factory 1 (VH)", font=('Segoe UI', 10, 'italic'), fg='#475569', bg='#F4F7FB')
    lbl_site_full.pack(side=tk.LEFT, padx=12)

    canvas_fact_kpi = tk.Canvas(tab_factory, bg='#F4F7FB', height=95, highlightthickness=0)
    canvas_fact_kpi.pack(fill=tk.X, padx=12, pady=(0, 6))

    fact_split = tk.PanedWindow(tab_factory, orient=tk.HORIZONTAL, bg='#E2E8F0', bd=0)
    fact_split.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 10))

    canvas_fact_charts = tk.Canvas(fact_split, bg='#F4F7FB', width=580, highlightthickness=0)
    fact_split.add(canvas_fact_charts)

    table_frame = tk.Frame(fact_split, bg='#FFFFFF', bd=1, relief=tk.SOLID)
    fact_split.add(table_frame)

    tk.Label(table_frame, text="4-STATION DEFECT MATRIX (TOP 5 MODELS)", font=('Segoe UI', 10, 'bold'), fg='#1F4E78', bg='#FFFFFF', pady=8).pack(fill=tk.X)
    
    cols_tree = ('Model', 'Station', 'Defect Issue', 'Qty')
    tree_defects = ttk.Treeview(table_frame, columns=cols_tree, show='headings', selectmode='browse')
    for c in cols_tree: tree_defects.heading(c, text=c)
    tree_defects.column('Model', width=180, anchor='w')
    tree_defects.column('Station', width=80, anchor='center')
    tree_defects.column('Defect Issue', width=180, anchor='w')
    tree_defects.column('Qty', width=60, anchor='e')

    tree_scroll = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=tree_defects.yview)
    tree_defects.configure(yscrollcommand=tree_scroll.set)
    tree_scroll.pack(side=tk.RIGHT, fill=tk.Y)
    tree_defects.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    def render_factory_tab(all_data):
        selected_site = site_combo.get()
        lbl_site_full.config(text=SITE_NAMES_FULL.get(selected_site, selected_site))
        
        canvas_fact_kpi.delete('all')
        canvas_fact_charts.delete('all')
        for item in tree_defects.get_children(): tree_defects.delete(item)

        if not all_data or selected_site not in all_data:
            return

        s_data = all_data[selected_site]
        top_c_name = s_data['top_countries'][0][0] if s_data['top_countries'] else "N/A"
        top_m_name = s_data['top_models'][0][0] if s_data['top_models'] else "N/A"
        top_d_name = s_data['top5_defects'][0][0] if s_data['top5_defects'] else "N/A"
        
        w_kpi = 270
        gap = 16
        draw_card(canvas_fact_kpi, 10 + 0 * (w_kpi + gap), 10, w_kpi, 75, f"{selected_site} QC Failures", f"{s_data['qc_total_fail']:,}", f"MA: {s_data['ma_total_fail']:,} Re-audits", '#1F4E78')
        draw_card(canvas_fact_kpi, 10 + 1 * (w_kpi + gap), 10, w_kpi, 75, "Top Fail Country", top_c_name, f"Qty: {s_data['top_countries'][0][1]}" if s_data['top_countries'] else "", '#2563EB')
        draw_card(canvas_fact_kpi, 10 + 2 * (w_kpi + gap), 10, w_kpi, 75, "Top Fail Model", f"{top_m_name[:18]}...", f"Qty: {s_data['top_models'][0][1]}" if s_data['top_models'] else "", '#D97706')
        draw_card(canvas_fact_kpi, 10 + 3 * (w_kpi + gap), 10, w_kpi, 75, "Primary Defect", f"{top_d_name[:18]}...", f"Qty: {int(s_data['top5_defects'][0][1])}" if s_data['top5_defects'] else "", '#DC2626')

        m_names = [m[0] for m in s_data['top_models']]
        m_vals = [m[1] for m in s_data['top_models']]
        draw_bar_chart(canvas_fact_charts, 10, 10, 560, 240, f"Top Failed Models (QC) — {selected_site}", m_names, m_vals, ['#1E3A8A', '#1D4ED8', '#2563EB', '#3B82F6', '#60A5FA'])

        d_names = [d[0] for d in s_data['top5_defects']]
        d_vals = [d[1] for d in s_data['top5_defects']]
        cur_cmap = pptx_generator.load_color_template(color_tpl_var.get()) if os.path.exists(color_tpl_var.get()) else {}
        d_colors = [f"#{pptx_generator.get_defect_hex(d, cur_cmap)}" for d in d_names]
        draw_bar_chart(canvas_fact_charts, 10, 260, 560, 240, f"Top 5 Defects (QC+MA) — {selected_site}", d_names, d_vals, d_colors)

        stations = [
            ('Re-Ins QC', s_data['re_qc_top3']),
            ('Re-Ins MA', s_data['re_ma_top3']),
            ('FTT', s_data['ftt_top3']),
            ('HFPA', s_data['hfpa_top3'])
        ]
        for st_name, st_list in stations:
            for m_info in st_list:
                for def_name, def_q in m_info['defects']:
                    tree_defects.insert('', tk.END, values=(m_info['model'], st_name, def_name, f"{int(def_q):,}"))

    site_combo.bind('<<ComboboxSelected>>', lambda e: render_factory_tab(loaded_data['sites']))

    # TAB 3: REGIONAL SUMMARY (VN vs INDO)
    canvas_reg = tk.Canvas(tab_regional, bg='#F4F7FB', highlightthickness=0)
    canvas_reg.pack(fill=tk.BOTH, expand=True, padx=12, pady=12)

    def render_regional_tab(all_data):
        canvas_reg.delete('all')
        if not all_data: return

        vh = all_data.get('VH', {'qc_total_fail': 0, 'ma_total_fail': 0})
        vh2 = all_data.get('VH2', {'qc_total_fail': 0, 'ma_total_fail': 0})
        jv = all_data.get('JV', {'qc_total_fail': 0, 'ma_total_fail': 0})
        jv2 = all_data.get('JV2', {'qc_total_fail': 0, 'ma_total_fail': 0})

        vn_tot = vh['qc_total_fail'] + vh2['qc_total_fail']
        indo_tot = jv['qc_total_fail'] + jv2['qc_total_fail']

        canvas_reg.create_rectangle(10, 10, 1200, 50, fill='#1F4E78', outline='')
        canvas_reg.create_text(605, 30, text="EXECUTIVE REGIONAL COMPARISON DASHBOARD — VIETNAM vs INDONESIA", font=('Segoe UI', 12, 'bold'), fill='#FFFFFF')

        # Vietnam Table
        canvas_reg.create_rectangle(10, 70, 590, 270, fill='#FFFFFF', outline='#CBD5E1', width=1)
        canvas_reg.create_rectangle(10, 70, 590, 105, fill='#2F5597', outline='')
        canvas_reg.create_text(25, 87, text="VIETNAM FACTORIES BREAKDOWN (VN)", font=('Segoe UI', 10, 'bold'), fill='#FFFFFF', anchor='w')

        canvas_reg.create_text(40, 125, text="Metric", font=('Segoe UI', 9, 'bold'), fill='#1F4E78', anchor='w')
        canvas_reg.create_text(190, 125, text="VH Factory", font=('Segoe UI', 9, 'bold'), fill='#1F4E78', anchor='center')
        canvas_reg.create_text(330, 125, text="VH2 Factory", font=('Segoe UI', 9, 'bold'), fill='#1F4E78', anchor='center')
        canvas_reg.create_text(475, 125, text="Total Vietnam", font=('Segoe UI', 9, 'bold'), fill='#1F4E78', anchor='center')
        canvas_reg.create_line(20, 140, 580, 140, fill='#E2E8F0')

        canvas_reg.create_text(40, 165, text="QC Fail Qty", font=('Segoe UI', 9), fill='#334155', anchor='w')
        canvas_reg.create_text(190, 165, text=f"{vh['qc_total_fail']:,}", font=('Segoe UI', 10, 'bold'), fill='#0F172A', anchor='center')
        canvas_reg.create_text(330, 165, text=f"{vh2['qc_total_fail']:,}", font=('Segoe UI', 10, 'bold'), fill='#0F172A', anchor='center')
        canvas_reg.create_text(475, 165, text=f"{vn_tot:,}", font=('Segoe UI', 11, 'bold'), fill='#1E3A8A', anchor='center')
        canvas_reg.create_line(20, 185, 580, 185, fill='#E2E8F0')

        vh_pct = (vh['qc_total_fail'] / vn_tot * 100) if vn_tot else 0
        vh2_pct = (vh2['qc_total_fail'] / vn_tot * 100) if vn_tot else 0
        canvas_reg.create_text(40, 210, text="Internal Share (%)", font=('Segoe UI', 9), fill='#334155', anchor='w')
        canvas_reg.create_text(190, 210, text=f"{vh_pct:.1f}%", font=('Segoe UI', 9, 'bold'), fill='#2563EB', anchor='center')
        canvas_reg.create_text(330, 210, text=f"{vh2_pct:.1f}%", font=('Segoe UI', 9, 'bold'), fill='#2563EB', anchor='center')
        canvas_reg.create_text(475, 210, text="100.0%", font=('Segoe UI', 9, 'bold'), fill='#1E3A8A', anchor='center')

        canvas_reg.create_text(40, 245, text="MA Re-audits", font=('Segoe UI', 9), fill='#64748B', anchor='w')
        canvas_reg.create_text(190, 245, text=f"{vh['ma_total_fail']:,}", font=('Segoe UI', 9), fill='#64748B', anchor='center')
        canvas_reg.create_text(330, 245, text=f"{vh2['ma_total_fail']:,}", font=('Segoe UI', 9), fill='#64748B', anchor='center')
        canvas_reg.create_text(475, 245, text=f"{vh['ma_total_fail']+vh2['ma_total_fail']:,}", font=('Segoe UI', 9, 'bold'), fill='#475569', anchor='center')

        # Indonesia Table
        canvas_reg.create_rectangle(620, 70, 1200, 270, fill='#FFFFFF', outline='#CBD5E1', width=1)
        canvas_reg.create_rectangle(620, 70, 1200, 105, fill='#C2410C', outline='')
        canvas_reg.create_text(635, 87, text="INDONESIA FACTORIES BREAKDOWN (INDO)", font=('Segoe UI', 10, 'bold'), fill='#FFFFFF', anchor='w')

        canvas_reg.create_text(650, 125, text="Metric", font=('Segoe UI', 9, 'bold'), fill='#9A3412', anchor='w')
        canvas_reg.create_text(800, 125, text="JV Factory", font=('Segoe UI', 9, 'bold'), fill='#9A3412', anchor='center')
        canvas_reg.create_text(940, 125, text="JV2 Factory", font=('Segoe UI', 9, 'bold'), fill='#9A3412', anchor='center')
        canvas_reg.create_text(1085, 125, text="Total Indonesia", font=('Segoe UI', 9, 'bold'), fill='#9A3412', anchor='center')
        canvas_reg.create_line(630, 140, 1190, 140, fill='#E2E8F0')

        canvas_reg.create_text(650, 165, text="QC Fail Qty", font=('Segoe UI', 9), fill='#334155', anchor='w')
        canvas_reg.create_text(800, 165, text=f"{jv['qc_total_fail']:,}", font=('Segoe UI', 10, 'bold'), fill='#0F172A', anchor='center')
        canvas_reg.create_text(940, 165, text=f"{jv2['qc_total_fail']:,}", font=('Segoe UI', 10, 'bold'), fill='#0F172A', anchor='center')
        canvas_reg.create_text(1085, 165, text=f"{indo_tot:,}", font=('Segoe UI', 11, 'bold'), fill='#9A3412', anchor='center')
        canvas_reg.create_line(630, 185, 1190, 185, fill='#E2E8F0')

        jv_pct = (jv['qc_total_fail'] / indo_tot * 100) if indo_tot else 0
        jv2_pct = (jv2['qc_total_fail'] / indo_tot * 100) if indo_tot else 0
        canvas_reg.create_text(650, 210, text="Internal Share (%)", font=('Segoe UI', 9), fill='#334155', anchor='w')
        canvas_reg.create_text(800, 210, text=f"{jv_pct:.1f}%", font=('Segoe UI', 9, 'bold'), fill='#EA580C', anchor='center')
        canvas_reg.create_text(940, 210, text=f"{jv2_pct:.1f}%", font=('Segoe UI', 9, 'bold'), fill='#EA580C', anchor='center')
        canvas_reg.create_text(1085, 210, text="100.0%", font=('Segoe UI', 9, 'bold'), fill='#9A3412', anchor='center')

        canvas_reg.create_text(650, 245, text="MA Re-audits", font=('Segoe UI', 9), fill='#64748B', anchor='w')
        canvas_reg.create_text(800, 245, text=f"{jv['ma_total_fail']:,}", font=('Segoe UI', 9), fill='#64748B', anchor='center')
        canvas_reg.create_text(940, 245, text=f"{jv2['ma_total_fail']:,}", font=('Segoe UI', 9), fill='#64748B', anchor='center')
        canvas_reg.create_text(1085, 245, text=f"{jv['ma_total_fail']+jv2['ma_total_fail']:,}", font=('Segoe UI', 9, 'bold'), fill='#475569', anchor='center')

        draw_bar_chart(canvas_reg, 10, 290, 590, 260, "4-Factory Audit Volume Distribution", ['VH (Vietnam)', 'VH2 (Vietnam)', 'JV (Indonesia)', 'JV2 (Indonesia)'], [vh['qc_total_fail'], vh2['qc_total_fail'], jv['qc_total_fail'], jv2['qc_total_fail']], ['#1F4E78', '#3B82F6', '#EA580C', '#F59E0B'])

        draw_donut_chart(canvas_reg, 910, 410, 85, "Regional Share of Total QC Failures", ['Vietnam (VN)', 'Indonesia (INDO)'], [vn_tot, indo_tot], ['#1F4E78', '#EA580C'])

    # TAB 4: CROSS-STATION ANALYSIS
    cross_top = tk.Frame(tab_cross, bg='#F4F7FB', pady=8)
    cross_top.pack(fill=tk.X, padx=12)

    tk.Label(cross_top, text="Station Comparison:", font=('Segoe UI', 10, 'bold'), fg='#1F4E78', bg='#F4F7FB').pack(side=tk.LEFT, padx=(0, 8))
    mode_cross_var = tk.StringVar(value="QC (FTT vs Re-Ins Initial Audit)")
    cb_cross = ttk.Combobox(cross_top, textvariable=mode_cross_var, values=["QC (FTT vs Re-Ins Initial Audit)", "MA (HFPA vs Re-Ins Re-Audit)"], state='readonly', font=('Segoe UI', 9), width=35)
    cb_cross.pack(side=tk.LEFT, padx=4)

    cross_table_frame = tk.Frame(tab_cross, bg='#FFFFFF', bd=1, relief=tk.SOLID)
    cross_table_frame.pack(fill=tk.BOTH, expand=True, padx=12, pady=10)

    cols_cross = ('Site', 'Audit Station', 'Shoe Model', 'Defect Description', 'Defect Qty')
    tree_cross = ttk.Treeview(cross_table_frame, columns=cols_cross, show='headings', selectmode='extended')
    for c in cols_cross: tree_cross.heading(c, text=c)
    tree_cross.column('Site', width=70, anchor='center')
    tree_cross.column('Audit Station', width=110, anchor='center')
    tree_cross.column('Shoe Model', width=220, anchor='w')
    tree_cross.column('Defect Description', width=280, anchor='w')
    tree_cross.column('Defect Qty', width=90, anchor='e')

    tree_cross_scroll = ttk.Scrollbar(cross_table_frame, orient=tk.VERTICAL, command=tree_cross.yview)
    tree_cross.configure(yscrollcommand=tree_cross_scroll.set)
    tree_cross_scroll.pack(side=tk.RIGHT, fill=tk.Y)
    tree_cross.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    def render_cross_tab(all_data):
        for item in tree_cross.get_children(): tree_cross.delete(item)
        if not all_data: return

        is_qc = "QC" in mode_cross_var.get()
        st1_name = "FTT" if is_qc else "HFPA"
        st2_name = "Re-Ins QC" if is_qc else "Re-Ins MA"

        for site in SITES:
            if site not in all_data: continue
            d = all_data[site]
            st1_list = d['ftt_top3'] if is_qc else d['hfpa_top3']
            st2_list = d['re_qc_top3'] if is_qc else d['re_ma_top3']

            for m_info in st1_list:
                for def_name, def_q in m_info['defects']:
                    tree_cross.insert('', tk.END, values=(site, st1_name, m_info['model'], def_name, f"{int(def_q):,}"))
            for m_info in st2_list:
                for def_name, def_q in m_info['defects']:
                    tree_cross.insert('', tk.END, values=(site, st2_name, m_info['model'], def_name, f"{int(def_q):,}"))

    cb_cross.bind('<<ComboboxSelected>>', lambda e: render_cross_tab(loaded_data['sites']))

    # TAB 5: EXPORTER & LOGS
    exp_top = tk.Frame(tab_logs, bg='#F4F7FB', pady=10)
    exp_top.pack(fill=tk.X, padx=12)

    f_paths = tk.LabelFrame(exp_top, text="Export File Destinations", font=('Segoe UI', 9, 'bold'), fg='#1F4E78', bg='#F4F7FB', padx=10, pady=8)
    f_paths.pack(fill=tk.X)

    make_ribbon_picker(f_paths, "Master Database Output (.xlsx):", db_out_var, "Select Database Output")
    make_ribbon_picker(f_paths, "Recycle Report Output (.xlsx):", recycle_out_var, "Select Recycle Report Output")
    make_ribbon_picker(f_paths, "PowerPoint Presentation (.pptx):", pptx_out_var, "Select PowerPoint Output")

    f_tpl = tk.Frame(exp_top, bg='#F4F7FB', pady=4)
    f_tpl.pack(fill=tk.X)
    tk.Label(f_tpl, text="PowerPoint Template:", font=('Segoe UI', 8, 'bold'), fg='#2F5597', bg='#F4F7FB').pack(side=tk.LEFT, padx=4)
    e_tpl = tk.Entry(f_tpl, textvariable=tpl_var, font=('Segoe UI', 9), bg='#FFFFFF', relief=tk.SOLID, bd=1)
    e_tpl.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=4)
    b_tpl = tk.Button(f_tpl, text="Browse...", command=lambda: browse_file(tpl_var, "Select Template PPTX", ".pptx", [("PowerPoint (*.pptx)", "*.pptx")]), width=10, bg='#E2E8F0', font=('Segoe UI', 8))
    b_tpl.pack(side=tk.RIGHT, padx=4)

    f_color_tpl = tk.Frame(exp_top, bg='#F4F7FB', pady=4)
    f_color_tpl.pack(fill=tk.X)
    tk.Label(f_color_tpl, text="Color Template (Defect Colors):", font=('Segoe UI', 8, 'bold'), fg='#2F5597', bg='#F4F7FB').pack(side=tk.LEFT, padx=4)
    e_ctpl = tk.Entry(f_color_tpl, textvariable=color_tpl_var, font=('Segoe UI', 9), bg='#FFFFFF', relief=tk.SOLID, bd=1)
    e_ctpl.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=4)
    b_ctpl = tk.Button(f_color_tpl, text="Browse...", command=lambda: browse_file(color_tpl_var, "Select Color Template", ".xlsx", [("Excel (*.xlsx)", "*.xlsx")]), width=10, bg='#E2E8F0', font=('Segoe UI', 8))
    b_ctpl.pack(side=tk.RIGHT, padx=4)

    f_actions = tk.Frame(exp_top, bg='#F4F7FB', pady=8)
    f_actions.pack(fill=tk.X)

    btn_t5_db = tk.Button(f_actions, text="📘 Export Database (.xlsx)", bg='#1F4E78', fg='white', font=('Segoe UI', 10, 'bold'), padx=10, pady=6)
    btn_t5_db.pack(side=tk.LEFT, padx=4)

    btn_t5_recycle = tk.Button(f_actions, text="📗 Export Recycle Report (.xlsx)", bg='#107C41', fg='white', font=('Segoe UI', 10, 'bold'), padx=10, pady=6)
    btn_t5_recycle.pack(side=tk.LEFT, padx=4)

    btn_t5_pptx = tk.Button(f_actions, text="📊 Export PowerPoint (.pptx)", bg='#D9531E', fg='white', font=('Segoe UI', 10, 'bold'), padx=12, pady=6)
    btn_t5_pptx.pack(side=tk.LEFT, padx=4)

    btn_t5_both = tk.Button(f_actions, text="🚀 Export All (Excel + PPTX)", bg='#0F2840', fg='white', font=('Segoe UI', 10, 'bold'), padx=12, pady=6)
    btn_t5_both.pack(side=tk.LEFT, padx=4)

    log_frame = tk.LabelFrame(tab_logs, text="Live Pipeline Console Log", font=('Segoe UI', 9, 'bold'), fg='#1F4E78', bg='#F4F7FB', padx=10, pady=8)
    log_frame.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 10))

    txt_logs = tk.Text(log_frame, bg='#0F172A', fg='#F8FAFC', font=('Consolas', 9), insertbackground='#FFFFFF', relief=tk.SOLID, bd=1)
    log_scroll = ttk.Scrollbar(log_frame, orient=tk.VERTICAL, command=txt_logs.yview)
    txt_logs.configure(yscrollcommand=log_scroll.set)
    log_scroll.pack(side=tk.RIGHT, fill=tk.Y)
    txt_logs.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    def log_ui(msg):
        txt_logs.insert(tk.END, f"[{datetime.datetime.now().strftime('%H:%M:%S')}] {msg}\n")
        txt_logs.see(tk.END)
        status_msg_var.set(msg[:60])
        root.update_idletasks()

    # Async Threaded Handlers
    def do_reload():
        btn_reload.config(state=tk.DISABLED)
        log_ui("Scanning folders and processing data across all sites dynamically...")
        
        def run():
            try:
                re_d = re_var.get().strip()
                ftt_d = ftt_var.get().strip()
                hfpa_d = hfpa_var.get().strip()
                
                period = detect_reporting_period(re_d)
                data = load_all_sites_data(re_d, ftt_d, hfpa_d, log_fn=log_ui)
                
                loaded_data['sites'] = data
                loaded_data['period'] = period
                
                def update_views():
                    period_var.set(f"Period: {period}")
                    default_db_name = format_db_filename(period)
                    default_recycle_name = format_recycle_filename(period)
                    default_pptx_name = format_pptx_filename(period)
                    db_out_var.set(os.path.join(cur_dir, default_db_name))
                    recycle_out_var.set(os.path.join(cur_dir, default_recycle_name))
                    pptx_out_var.set(os.path.join(cur_dir, default_pptx_name))
                    
                    render_overview(data)
                    render_factory_tab(data)
                    render_regional_tab(data)
                    render_cross_tab(data)
                    log_ui(f"✓ Data loaded and visualized successfully! (Period: {period})")
                    btn_reload.config(state=tk.NORMAL)
                    
                root.after(0, update_views)
            except Exception as e:
                def err_ui():
                    log_ui(f"[ERROR] Failed to load data: {e}")
                    messagebox.showerror("Data Load Error", str(e))
                    btn_reload.config(state=tk.NORMAL)
                root.after(0, err_ui)

        threading.Thread(target=run, daemon=True).start()

    def do_gen_database():
        def run():
            try:
                re_d = re_var.get().strip(); ftt_d = ftt_var.get().strip(); hfpa_d = hfpa_var.get().strip()
                out_f = db_out_var.get().strip()
                generate_database_workbook(re_d, ftt_d, hfpa_d, out_f, log_fn=log_ui)
                messagebox.showinfo("Success", f"Re-Inspection Database generated successfully!\n\nSaved to: {out_f}")
            except Exception as e:
                log_ui(f"[ERROR] Database generation failed: {e}")
                messagebox.showerror("Generation Error", str(e))

        threading.Thread(target=run, daemon=True).start()

    def do_gen_recycle():
        def run():
            try:
                re_d = re_var.get().strip(); ftt_d = ftt_var.get().strip(); hfpa_d = hfpa_var.get().strip()
                out_f = recycle_out_var.get().strip()
                generate_recycle_workbook(re_d, ftt_d, hfpa_d, out_f, log_fn=log_ui)
                messagebox.showinfo("Success", f"Recycle Report generated successfully!\n\nSaved to: {out_f}")
            except Exception as e:
                log_ui(f"[ERROR] Recycle Report generation failed: {e}")
                messagebox.showerror("Generation Error", str(e))

        threading.Thread(target=run, daemon=True).start()

    def do_gen_pptx():
        def run():
            try:
                re_d = re_var.get().strip(); ftt_d = ftt_var.get().strip(); hfpa_d = hfpa_var.get().strip()
                tpl_f = tpl_var.get().strip()
                color_f = color_tpl_var.get().strip()
                period = detect_reporting_period(re_d)
                out_p = pptx_out_var.get().strip()
                if not out_p:
                    out_p = os.path.join(cur_dir, format_pptx_filename(period))
                generate_presentation_report(re_d, ftt_d, hfpa_d, tpl_f, out_p, month_label=period, color_template=color_f, log_fn=log_ui)
                messagebox.showinfo("Success", f"PowerPoint Presentation generated successfully!\n\nSaved to: {out_p}")
            except Exception as e:
                log_ui(f"[ERROR] PowerPoint export failed: {e}")
                messagebox.showerror("PowerPoint Export Error", str(e))

        threading.Thread(target=run, daemon=True).start()

    def do_gen_both():
        def run():
            try:
                re_d = re_var.get().strip(); ftt_d = ftt_var.get().strip(); hfpa_d = hfpa_var.get().strip()
                db_out = db_out_var.get().strip()
                rec_out = recycle_out_var.get().strip()
                period = detect_reporting_period(re_d)
                pptx_out = pptx_out_var.get().strip()
                if not pptx_out:
                    pptx_out = os.path.join(cur_dir, format_pptx_filename(period))
                tpl_f = tpl_var.get().strip()
                color_f = color_tpl_var.get().strip()
                generate_all_reports(re_d, ftt_d, hfpa_d, db_out, rec_out, pptx_out, tpl_f, color_template=color_f, log_fn=log_ui)
                messagebox.showinfo("Success", f"All reports generated successfully!\n\n1. Database: {db_out}\n2. Recycle: {rec_out}\n3. Presentation: {pptx_out}")
            except Exception as e:
                log_ui(f"[ERROR] Batch generation failed: {e}")
                messagebox.showerror("Generation Error", str(e))

        threading.Thread(target=run, daemon=True).start()

    btn_reload.config(command=do_reload)
    btn_gen_db.config(command=do_gen_database)
    btn_gen_recycle.config(command=do_gen_recycle)
    btn_gen_pptx.config(command=do_gen_pptx)
    btn_gen_both.config(command=do_gen_both)

    btn_t5_db.config(command=do_gen_database)
    btn_t5_recycle.config(command=do_gen_recycle)
    btn_t5_pptx.config(command=do_gen_pptx)
    btn_t5_both.config(command=do_gen_both)

    log_ui("Initializing Re-Inspection Intelligence Suite...")
    root.after(100, do_reload)

    root.mainloop()


# ==============================================================================
#  CLI & MAIN ENTRYPOINT
# ==============================================================================

if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description="Unified Re-Inspection & Recycle Report Suite with Python GUI Dashboard & PPTX Export"
    )
    parser.add_argument('--gui', action='store_true', help='Launch interactive native Python GUI Dashboard')
    parser.add_argument('--mode', choices=['database', 'recycle', 'pptx', 'all'], default='gui', help='Generation mode (default: gui)')
    parser.add_argument('--re-dir', default='Re-Inspection', help='Path to Re-Inspection folder')
    parser.add_argument('--ftt-dir', default='FTT', help='Path to FTT folder')
    parser.add_argument('--hfpa-dir', default='HFPA', help='Path to HFPA folder')
    parser.add_argument('--template-pptx', default='RE-INS REPORT.APR.2026.pptx', help='Template PowerPoint file path')
    parser.add_argument('--color-template', default='Color_Template.xlsx', help='Path to Color_Template.xlsx')
    parser.add_argument('--output-db', default=None, help='Output path for Master Database (defaults to Re-Inspection-Database-Month-Year.xlsx)')
    parser.add_argument('--output-recycle', default=None, help='Output path for Recycle Report (defaults to Recycle_Report-Month-Year.xlsx)')
    parser.add_argument('--output-pptx', default=None, help='Output path for PowerPoint Presentation (defaults to Re-Ins Report-Month-Year.pptx)')
    parser.add_argument('--output', default=None, help='Generic output path override')
    parser.add_argument('--month', default=None, help='Month label override (e.g. "June 2026")')

    args = parser.parse_args()

    if args.gui or (args.mode == 'gui' and len(sys.argv) == 1):
        launch_dashboard(args.re_dir, args.ftt_dir, args.hfpa_dir, args.template_pptx, args.color_template)
    elif args.mode == 'database':
        period = args.month if args.month else detect_reporting_period(args.re_dir)
        out_f = args.output if args.output else (args.output_db if args.output_db else format_db_filename(period))
        generate_database_workbook(args.re_dir, args.ftt_dir, args.hfpa_dir, out_f, month_label=period, color_template=args.color_template)
    elif args.mode == 'recycle':
        period = args.month if args.month else detect_reporting_period(args.re_dir)
        out_f = args.output if args.output else (args.output_recycle if args.output_recycle else format_recycle_filename(period))
        generate_recycle_workbook(args.re_dir, args.ftt_dir, args.hfpa_dir, out_f, month_label=period)
    elif args.mode == 'pptx':
        period = args.month if args.month else detect_reporting_period(args.re_dir)
        out_p = args.output if args.output else (args.output_pptx if args.output_pptx else format_pptx_filename(period))
        generate_presentation_report(args.re_dir, args.ftt_dir, args.hfpa_dir, args.template_pptx, out_p, month_label=period, color_template=args.color_template)
    elif args.mode == 'all':
        period = args.month if args.month else detect_reporting_period(args.re_dir)
        db_p = args.output_db if args.output_db else format_db_filename(period)
        rec_p = args.output_recycle if args.output_recycle else format_recycle_filename(period)
        out_p = args.output_pptx if args.output_pptx else format_pptx_filename(period)
        generate_all_reports(args.re_dir, args.ftt_dir, args.hfpa_dir, db_p, rec_p, out_p, args.template_pptx, color_template=args.color_template)
    else:
        launch_dashboard(args.re_dir, args.ftt_dir, args.hfpa_dir, args.template_pptx, args.color_template)
