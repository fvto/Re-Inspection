"""
pptx_generator.py — PowerPoint Presentation Exporter for Re-Inspection Reports
=============================================================================
Replicates the exact layout, styling, and charts of 'RE-INS REPORT.APR.2026.pptx'
by populating fresh data from Re-Inspection, FTT, and HFPA data sources, with:
  1. Dynamic slide 2 KPIs, Top 5 model tables, Doughnut charts, Defect bar charts, Country 3D pie charts.
  2. Dynamic slides 3 & 4 4-Station Comparison 100% Stacked Column Charts (Charts 12 to 19).
  3. Strict defect color mapping from 'Color_Template.xlsx' across all comparison charts and legend tables.
"""

import os
import io
import re
import math
import copy
import datetime
import colorsys
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

NS_C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'

ET.register_namespace('c', NS_C)
ET.register_namespace('a', NS_A)
ET.register_namespace('p', NS_P)

# ==============================================================================
#  COLOR TEMPLATE LOADER, SYNCHRONIZER & UNIQUE COLOR GENERATOR
# ==============================================================================

def hex_to_rgb(hex_str):
    """Convert 6-character HEX string to (R, G, B) integer tuple."""
    hex_str = str(hex_str).strip().upper().replace('#', '').zfill(6)
    return int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)

def rgb_to_hex(r, g, b):
    """Convert (R, G, B) integers to uppercase 6-character HEX string."""
    return f"{int(r):02X}{int(g):02X}{int(b):02X}"

def color_distance(rgb1, rgb2):
    """Euclidean distance between two RGB colors."""
    return math.sqrt((rgb1[0] - rgb2[0])**2 + (rgb1[1] - rgb2[1])**2 + (rgb1[2] - rgb2[2])**2)

def generate_unique_color(existing_rgbs, existing_hexes, seed_index=0):
    """
    Generate a vibrant, corporate-ready unique color that does not duplicate or clash
    with existing colors in Color_Template.xlsx.
    Avoids stark white, stark black, and neutral grays.
    """
    golden_ratio = 0.618033988749895
    forbidden_rgbs = [
        (255, 255, 255), (0, 0, 0), (220, 220, 220), (240, 240, 240), (20, 20, 20)
    ]
    min_dist_threshold = 36.0

    for attempt in range(3000):
        h = (seed_index * golden_ratio + attempt * 0.137) % 1.0
        s = 0.55 + (attempt % 5) * 0.08
        v = 0.50 + ((attempt // 5) % 4) * 0.10
        r, g, b = [int(round(c * 255)) for c in colorsys.hsv_to_rgb(h, s, v)]
        cand_hex = rgb_to_hex(r, g, b)

        if cand_hex in existing_hexes:
            continue
        if any(color_distance((r, g, b), f) < 40 for f in forbidden_rgbs):
            continue
        curr_min = min((color_distance((r, g, b), ext) for ext in existing_rgbs), default=999)
        if curr_min >= min_dist_threshold:
            return cand_hex, r, g, b

    # Fallback with smaller distance
    for attempt in range(1000):
        h = (seed_index * golden_ratio + attempt * 0.213) % 1.0
        r, g, b = [int(round(c * 255)) for c in colorsys.hsv_to_rgb(h, 0.60, 0.60)]
        cand_hex = rgb_to_hex(r, g, b)
        if cand_hex not in existing_hexes and not any(color_distance((r, g, b), f) < 30 for f in forbidden_rgbs):
            return cand_hex, r, g, b

    return rgb_to_hex(r, g, b), r, g, b

def norm_defect_name(s):
    """Normalize defect name for fuzzy matching (removes whitespace, slashes, dashes, parentheses)."""
    return re.sub(r'[\s/&_\-\(\)]+', '', str(s).lower())

def sync_and_register_defects(defect_names=None, color_path="Color_Template.xlsx", log_fn=print):
    """
    Ensure Color_Template.xlsx is 100% compliant and rendered in both Microsoft Excel
    and WPS Office:
      - Cleans problematic conditional formatting rules that break WPS Office rendering.
      - Sets explicit PatternFill with both start_color and end_color in aRGB hex (FFxxxxxx).
      - Adds borders and standard alignments.
      - Auto-generates unique, non-duplicating colors for any new defects not present yet.
    Returns:
      Updated color_map dictionary (defect_name -> 6-char HEX).
    """
    if not os.path.exists(color_path):
        return {}

    try:
        wb = openpyxl.load_workbook(color_path)
    except Exception as e:
        log_fn(f"[WARNING] Could not open {color_path} for synchronization: {e}")
        return load_color_template(color_path)

    ws1 = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active

    # Clear problematic conditional formatting that breaks WPS Office
    if hasattr(ws1, 'conditional_formatting'):
        try:
            ws1.conditional_formatting._cf_rules.clear()
        except:
            pass

    existing_names = {}
    existing_norm_keys = {}
    existing_hexes = set()
    existing_rgbs = []

    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )
    font_def = Font(name='Calibri', size=10)

    # 1. Standardize Sheet1
    max_r = ws1.max_row
    for r in range(5, max_r + 1):
        d_val = ws1.cell(r, 3).value
        hex_val = ws1.cell(r, 8).value
        if d_val is not None and str(d_val).strip():
            d_name = str(d_val).strip().replace('\xa0', ' ')
            rv = ws1.cell(r, 5).value
            gv = ws1.cell(r, 6).value
            bv = ws1.cell(r, 7).value
            has_valid_rgb = all(isinstance(v, (int, float)) and 0 <= int(v) <= 255 for v in [rv, gv, bv] if v is not None) and all(v is not None for v in [rv, gv, bv])

            if hex_val is not None:
                clean_hex = str(hex_val).strip().upper().replace('#', '').zfill(6)
            elif has_valid_rgb:
                clean_hex = rgb_to_hex(int(rv), int(gv), int(bv))
            else:
                clean_hex = rgb_to_hex(rv or 128, gv or 128, bv or 128)

            rgb_t = hex_to_rgb(clean_hex)
            ws1.cell(r, 3, d_name)
            ws1.cell(r, 5, rgb_t[0]).alignment = Alignment(horizontal='right', vertical='center')
            ws1.cell(r, 6, rgb_t[1]).alignment = Alignment(horizontal='right', vertical='center')
            ws1.cell(r, 7, rgb_t[2]).alignment = Alignment(horizontal='right', vertical='center')
            ws1.cell(r, 8, clean_hex).alignment = Alignment(horizontal='center', vertical='center')

            # Dual-color PatternFill for WPS & Excel compatibility
            argb = 'FF' + clean_hex
            ws1.cell(r, 4).fill = PatternFill(fill_type='solid', start_color=argb, end_color=argb)
            ws1.cell(r, 4).border = thin_border
            ws1.cell(r, 4).value = None

            existing_names[d_name] = clean_hex
            existing_names[d_name.lower()] = clean_hex
            existing_norm_keys[norm_defect_name(d_name)] = clean_hex
            existing_hexes.add(clean_hex)
            existing_rgbs.append(rgb_t)

    # 2. Standardize sheet bc if present
    if 'bc' in wb.sheetnames:
        ws_bc = wb['bc']
        if hasattr(ws_bc, 'conditional_formatting'):
            try:
                ws_bc.conditional_formatting._cf_rules.clear()
            except:
                pass
        for r in range(2, ws_bc.max_row + 1):
            d_val = ws_bc.cell(r, 1).value
            if d_val is not None and str(d_val).strip():
                d_name = str(d_val).strip().replace('\xa0', ' ')
                hex_val = ws_bc.cell(r, 5).value
                rv = ws_bc.cell(r, 2).value or 128
                gv = ws_bc.cell(r, 3).value or 128
                bv = ws_bc.cell(r, 4).value or 128
                clean_hex = str(hex_val).strip().upper().replace('#', '').zfill(6) if hex_val else rgb_to_hex(rv, gv, bv)
                ws_bc.cell(r, 5, clean_hex)

                rgb_t = hex_to_rgb(clean_hex)
                argb = 'FF' + clean_hex
                ws_bc.cell(r, 6).fill = PatternFill(fill_type='solid', start_color=argb, end_color=argb)
                ws_bc.cell(r, 6).border = thin_border

                existing_names[d_name] = clean_hex
                existing_names[d_name.lower()] = clean_hex
                existing_norm_keys[norm_defect_name(d_name)] = clean_hex
                existing_hexes.add(clean_hex)
                existing_rgbs.append(rgb_t)

    # 3. Add any new unique defects
    added_count = 0
    if defect_names:
        cur_row = ws1.max_row + 1
        for idx, d_raw in enumerate(sorted(set(defect_names))):
            d_clean = str(d_raw).strip().replace('\xa0', ' ')
            if not d_clean or d_clean.lower() in existing_names:
                continue

            n_key = norm_defect_name(d_clean)
            if n_key in existing_norm_keys:
                matched_hex = existing_norm_keys[n_key]
                existing_names[d_clean] = matched_hex
                existing_names[d_clean.lower()] = matched_hex
                continue

            new_hex, nr, ng, nb = generate_unique_color(existing_rgbs, existing_hexes, seed_index=idx + added_count)
            existing_hexes.add(new_hex)
            existing_rgbs.append((nr, ng, nb))
            existing_names[d_clean] = new_hex
            existing_names[d_clean.lower()] = new_hex
            existing_norm_keys[n_key] = new_hex

            ws1.cell(cur_row, 3, d_clean).font = font_def
            ws1.cell(cur_row, 3).alignment = Alignment(horizontal='left', vertical='center')

            argb = 'FF' + new_hex
            ws1.cell(cur_row, 4).fill = PatternFill(fill_type='solid', start_color=argb, end_color=argb)
            ws1.cell(cur_row, 4).border = thin_border
            ws1.cell(cur_row, 4).value = None

            ws1.cell(cur_row, 5, nr).alignment = Alignment(horizontal='right', vertical='center')
            ws1.cell(cur_row, 5).font = font_def
            ws1.cell(cur_row, 6, ng).alignment = Alignment(horizontal='right', vertical='center')
            ws1.cell(cur_row, 6).font = font_def
            ws1.cell(cur_row, 7, nb).alignment = Alignment(horizontal='right', vertical='center')
            ws1.cell(cur_row, 7).font = font_def
            ws1.cell(cur_row, 8, new_hex).alignment = Alignment(horizontal='center', vertical='center')
            ws1.cell(cur_row, 8).font = font_def

            ws1.row_dimensions[cur_row].height = 20
            cur_row += 1
            added_count += 1

    try:
        wb.save(color_path)
        if added_count > 0:
            log_fn(f"[OK] Registered {added_count} new defect colors into {color_path} (WPS & Excel compatible).")
    except Exception as e:
        log_fn(f"[WARNING] Could not save updated {color_path}: {e}")

    return existing_names

def load_color_template(color_path="Color_Template.xlsx"):
    """Load defect type to HEX color map from Color_Template.xlsx (reads both Sheet1 and bc)."""
    color_map = {}
    if os.path.exists(color_path):
        try:
            wb = openpyxl.load_workbook(color_path, data_only=True)
            if 'Sheet1' in wb.sheetnames:
                ws1 = wb['Sheet1']
                for r in range(5, ws1.max_row + 1):
                    dtype = ws1.cell(r, 3).value
                    hex_val = ws1.cell(r, 8).value
                    if dtype is not None and hex_val is not None:
                        clean_name = str(dtype).strip().replace('\xa0', ' ')
                        clean_hex = str(hex_val).strip().upper().replace('#', '').zfill(6)
                        color_map[clean_name] = clean_hex
                        color_map[clean_name.lower()] = clean_hex

            if 'bc' in wb.sheetnames:
                ws_bc = wb['bc']
                for r in range(2, ws_bc.max_row + 1):
                    dtype = ws_bc.cell(r, 1).value
                    hex_val = ws_bc.cell(r, 5).value
                    if dtype is not None and hex_val is not None:
                        clean_name = str(dtype).strip().replace('\xa0', ' ')
                        clean_hex = str(hex_val).strip().upper().replace('#', '').zfill(6)
                        color_map[clean_name] = clean_hex
                        color_map[clean_name.lower()] = clean_hex
        except Exception as e:
            print(f"[WARNING] Could not read {color_path}: {e}")
    return color_map

def get_defect_hex(defect_name, color_map):
    """Retrieve exact HEX color for defect name with intelligent fallback."""
    if not defect_name:
        return 'DCDCDC'
    d_clean = str(defect_name).strip().replace('\xa0', ' ')
    if d_clean in color_map:
        return color_map[d_clean]
    if d_clean.lower() in color_map:
        return color_map[d_clean.lower()]
    norm_target = norm_defect_name(d_clean)
    for k, v in color_map.items():
        if norm_defect_name(k) == norm_target:
            return v
    return 'DCDCDC'



# ==============================================================================
#  OPENXML CHART UPDATERS (SLIDE 2)
# ==============================================================================

def _find_cat_container(ser):
    """Find the category data container (strCache or strData) inside a series."""
    cat_container = ser.find(f'.//{{{NS_C}}}cat//{{{NS_C}}}strCache')
    if cat_container is None:
        cat_container = ser.find(f'.//{{{NS_C}}}cat//{{{NS_C}}}strData')
    return cat_container

def _find_val_container(ser):
    """Find the value data container (numCache or numData) inside a series."""
    val_container = ser.find(f'.//{{{NS_C}}}val//{{{NS_C}}}numCache')
    if val_container is None:
        val_container = ser.find(f'.//{{{NS_C}}}val//{{{NS_C}}}numData')
    return val_container

def _update_cat_points(container, categories):
    """Replace category points in a strCache/strData container."""
    if container is None: return
    for pt in list(container.findall(f'{{{NS_C}}}pt')): container.remove(pt)
    cnt = container.find(f'{{{NS_C}}}ptCount')
    if cnt is not None: cnt.set('val', str(len(categories)))
    for idx, cat_name in enumerate(categories):
        pt = ET.SubElement(container, f'{{{NS_C}}}pt', idx=str(idx))
        ET.SubElement(pt, f'{{{NS_C}}}v').text = str(cat_name)

def _update_val_points(container, values, is_pct=False):
    """Replace value points in a numCache/numData container."""
    if container is None: return
    for pt in list(container.findall(f'{{{NS_C}}}pt')): container.remove(pt)
    cnt = container.find(f'{{{NS_C}}}ptCount')
    if cnt is not None: cnt.set('val', str(len(values)))
    for idx, val in enumerate(values):
        pt = ET.SubElement(container, f'{{{NS_C}}}pt', idx=str(idx))
        if is_pct:
            v_text = f"{val:.16E}" if 0 < val < 0.1 else f"{val:.16f}".rstrip('0').rstrip('.')
        else:
            v_text = str(int(val))
        ET.SubElement(pt, f'{{{NS_C}}}v').text = v_text

def _update_doughnut_chart(xml_content, title, categories, values):
    """Update a 2-slice or N-slice doughnut chart in chart XML."""
    root = ET.fromstring(xml_content)
    ser = root.find(f'.//{{{NS_C}}}ser')
    if ser is None: return xml_content

    tx_v = ser.find(f'.//{{{NS_C}}}tx//{{{NS_C}}}v')
    if tx_v is not None: tx_v.text = str(title)

    _update_cat_points(_find_cat_container(ser), categories)
    _update_val_points(_find_val_container(ser), values)

    return ET.tostring(root, encoding='utf-8', xml_declaration=True)

def _update_bar_chart_dLbls(ser, categories, values2):
    """
    Update data labels (<c:dLbls>) on the % Defect series in Slide 2 horizontal bar charts.
    Formats each label matching the reference layout:
      Run 1: "<Defect Name>, " in bold 9pt text.
      Run 2: "<Percentage>" in bold 9pt text with green color #00B050 (e.g. 15.3%, 22%).
    """
    dLbls = ser.find(f'{{{NS_C}}}dLbls')
    if dLbls is None:
        dLbls = ET.SubElement(ser, f'{{{NS_C}}}dLbls')

    # Remove extra dLbl elements beyond len(categories)
    existing_dlbls = {}
    for dl in list(dLbls.findall(f'{{{NS_C}}}dLbl')):
        idx_el = dl.find(f'{{{NS_C}}}idx')
        if idx_el is not None:
            try:
                idx_val = int(idx_el.get('val', '-1'))
                if idx_val >= len(categories) or idx_val < 0:
                    dLbls.remove(dl)
                else:
                    existing_dlbls[idx_val] = dl
            except ValueError:
                dLbls.remove(dl)

    for i, cat_name in enumerate(categories):
        pct_val = values2[i] if i < len(values2) else 0.0
        pct_round = round(pct_val * 100, 1)
        pct_str = f"{int(pct_round)}%" if pct_round.is_integer() else f"{pct_round:.1f}%"

        dl = existing_dlbls.get(i)
        if dl is None:
            dl = ET.SubElement(dLbls, f'{{{NS_C}}}dLbl')
            ET.SubElement(dl, f'{{{NS_C}}}idx', val=str(i))

        pos = dl.find(f'{{{NS_C}}}dLblPos')
        if pos is None:
            ET.SubElement(dl, f'{{{NS_C}}}dLblPos', val="outEnd")
        else:
            pos.set('val', 'outEnd')

        for tag, val in [('showLegendKey', '0'), ('showVal', '1'), ('showCatName', '1'), ('showSerName', '0'), ('showPercent', '0'), ('showBubbleSize', '0')]:
            el = dl.find(f'{{{NS_C}}}{tag}')
            if el is None:
                ET.SubElement(dl, f'{{{NS_C}}}{tag}', val=val)
            else:
                el.set('val', val)

        tx = dl.find(f'{{{NS_C}}}tx')
        if tx is None:
            tx = ET.SubElement(dl, f'{{{NS_C}}}tx')
        rich = tx.find(f'{{{NS_C}}}rich')
        if rich is None:
            rich = ET.SubElement(tx, f'{{{NS_C}}}rich')
            ET.SubElement(rich, f'{{{NS_A}}}bodyPr', rot="0", spcFirstLastPara="0", vertOverflow="ellipsis", vert="horz", wrap="square", lIns="38100", tIns="19050", rIns="38100", bIns="19050", anchor="ctr", anchorCtr="0")
            ET.SubElement(rich, f'{{{NS_A}}}lstStyle')

        p = rich.find(f'{{{NS_A}}}p')
        if p is None:
            p = ET.SubElement(rich, f'{{{NS_A}}}p')

        # Remove existing runs and endParaRPr
        for child in list(p):
            if child.tag in [f'{{{NS_A}}}r', f'{{{NS_A}}}endParaRPr']:
                p.remove(child)

        # Run 1: Category Name + ", "
        r1 = ET.SubElement(p, f'{{{NS_A}}}r')
        r1_pr = ET.SubElement(r1, f'{{{NS_A}}}rPr', lang="en-US", sz="900", b="1", dirty="0")
        r1_fill = ET.SubElement(r1_pr, f'{{{NS_A}}}solidFill')
        ET.SubElement(r1_fill, f'{{{NS_A}}}srgbClr', val="000000")
        ET.SubElement(r1, f'{{{NS_A}}}t').text = f"{cat_name}, "

        # Run 2: Percentage in Green (#00B050)
        r2 = ET.SubElement(p, f'{{{NS_A}}}r')
        r2_pr = ET.SubElement(r2, f'{{{NS_A}}}rPr', lang="en-US", sz="900", b="1", dirty="0")
        r2_fill = ET.SubElement(r2_pr, f'{{{NS_A}}}solidFill')
        ET.SubElement(r2_fill, f'{{{NS_A}}}srgbClr', val="00B050")
        ET.SubElement(r2, f'{{{NS_A}}}t').text = pct_str

        # End paragraph properties with matching green
        end_pr = ET.SubElement(p, f'{{{NS_A}}}endParaRPr', lang="en-US", sz="900", b="1", dirty="0")
        end_fill = ET.SubElement(end_pr, f'{{{NS_A}}}solidFill')
        ET.SubElement(end_fill, f'{{{NS_A}}}srgbClr', val="00B050")

def _update_bar_chart_2series(xml_content, categories, values1, values2):
    """Update a 2-series bar chart (Quantity and Percentage) in chart XML."""
    root = ET.fromstring(xml_content)
    series_list = root.findall(f'.//{{{NS_C}}}ser')
    if not series_list: return xml_content

    for s_idx, (vals, is_pct) in enumerate([(values1, False), (values2, True)]):
        if s_idx >= len(series_list): break
        ser = series_list[s_idx]
        _update_cat_points(_find_cat_container(ser), categories)
        _update_val_points(_find_val_container(ser), vals, is_pct=is_pct)
        if is_pct:
            _update_bar_chart_dLbls(ser, categories, values2)

    # --- SAFE ZONE ENFORCEMENT ---
    # Dynamically scale valAx max so the longest bar occupies at most 44-48% of the axis width,
    # reserving the remaining 52%+ as a guaranteed clean safe zone for category names and percentages.
    valAx = root.find(f'.//{{{NS_C}}}valAx')
    if valAx is not None and values1:
        valid_vals = [float(v) for v in values1 if v is not None and float(v) > 0]
        if valid_vals:
            max_val = max(valid_vals)
            target_max = max_val * 2.22
            if target_max <= 60:
                scale_max = int(math.ceil(target_max / 10.0) * 10)
                major_unit = max(10, scale_max // 4)
            elif target_max <= 150:
                scale_max = int(math.ceil(target_max / 25.0) * 25)
                major_unit = 25
            elif target_max <= 350:
                scale_max = int(math.ceil(target_max / 50.0) * 50)
                major_unit = 50
            elif target_max <= 750:
                scale_max = int(math.ceil(target_max / 100.0) * 100)
                major_unit = 100
            elif target_max <= 1500:
                scale_max = int(math.ceil(target_max / 200.0) * 200)
                major_unit = 200
            else:
                scale_max = int(math.ceil(target_max / 500.0) * 500)
                major_unit = 500

            scaling = valAx.find(f'{{{NS_C}}}scaling')
            if scaling is not None:
                max_el = scaling.find(f'{{{NS_C}}}max')
                if max_el is not None:
                    max_el.set('val', str(scale_max))
                else:
                    ET.SubElement(scaling, f'{{{NS_C}}}max', val=str(scale_max))
            mu_el = valAx.find(f'{{{NS_C}}}majorUnit')
            if mu_el is not None:
                mu_el.set('val', str(major_unit))
            else:
                ET.SubElement(valAx, f'{{{NS_C}}}majorUnit', val=str(major_unit))

    return ET.tostring(root, encoding='utf-8', xml_declaration=True)

def _update_pie3d_chart(xml_content, categories, values1, values2):
    """
    Update a 3D Pie Chart with categories and Grand Total + Percentage series.
    Enforces clean, high-contrast, non-overlapping typography without truncation.
    """
    root = ET.fromstring(xml_content)
    pie = root.find(f'.//{{{NS_C}}}pie3DChart')
    if pie is None:
        pie = root.find(f'.//{{{NS_C}}}pieChart')
    if pie is None: return xml_content

    series_list = pie.findall(f'{{{NS_C}}}ser')
    for s_idx, (vals, is_pct) in enumerate([(values1, False), (values2, True)]):
        if s_idx >= len(series_list): break
        ser = series_list[s_idx]
        _update_cat_points(_find_cat_container(ser), categories)
        _update_val_points(_find_val_container(ser), vals, is_pct=is_pct)

    # Clean and optimize data labels on primary series (ser[0])
    if len(series_list) > 0:
        ser0 = series_list[0]
        dLbls0 = ser0.find(f'{{{NS_C}}}dLbls')
        if dLbls0 is None:
            dLbls0 = ET.SubElement(ser0, f'{{{NS_C}}}dLbls')

        # Wipe out any legacy overrides with small padding or ellipsis truncation
        for child in list(dLbls0):
            dLbls0.remove(child)

        # Build clean data label configuration
        ET.SubElement(dLbls0, f'{{{NS_C}}}numFmt', formatCode='0.0%', sourceLinked='0')
        ET.SubElement(dLbls0, f'{{{NS_C}}}dLblPos', val='bestFit')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showLegendKey', val='0')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showVal', val='0')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showCatName', val='0')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showSerName', val='0')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showPercent', val='1')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showBubbleSize', val='0')
        ET.SubElement(dLbls0, f'{{{NS_C}}}showLeaderLines', val='1')

        # Charcoal bold text (sz=800, 8pt) with overflow enabled to eliminate '40...'
        txPr = ET.SubElement(dLbls0, f'{{{NS_C}}}txPr')
        ET.SubElement(txPr, f'{{{NS_A}}}bodyPr', rot='0', spcFirstLastPara='0', vertOverflow='overflow', horzOverflow='overflow', vert='horz', wrap='none', lIns='0', tIns='0', rIns='0', bIns='0', anchor='ctr', anchorCtr='1')
        ET.SubElement(txPr, f'{{{NS_A}}}lstStyle')
        p = ET.SubElement(txPr, f'{{{NS_A}}}p')
        pPr = ET.SubElement(p, f'{{{NS_A}}}pPr')
        defRPr = ET.SubElement(pPr, f'{{{NS_A}}}defRPr', lang='en-US', sz='800', b='1', i='0', u='none', strike='noStrike', kern='1200', baseline='0')
        sf = ET.SubElement(defRPr, f'{{{NS_A}}}solidFill')
        ET.SubElement(sf, f'{{{NS_A}}}srgbClr', val='1F2937')
        ET.SubElement(p, f'{{{NS_A}}}endParaRPr', lang='en-US', sz='800', b='1')

    # Delete data labels on secondary series (ser[1]) to prevent duplicate overlapping labels
    if len(series_list) > 1:
        ser1 = series_list[1]
        dLbls1 = ser1.find(f'{{{NS_C}}}dLbls')
        if dLbls1 is None:
            dLbls1 = ET.SubElement(ser1, f'{{{NS_C}}}dLbls')
        for child in list(dLbls1):
            dLbls1.remove(child)
        ET.SubElement(dLbls1, f'{{{NS_C}}}delete', val='1')

    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


# ==============================================================================
#  OPENXML COMPARISON CHART UPDATERS (SLIDES 3 & 4)
# ==============================================================================

def _update_comparison_stacked_chart(xml_content, st1_name, st2_name, models_alphabetical, st1_model_defects, st2_model_defects, color_map):
    """
    Update 100% Stacked Column Chart for Station Cross Comparisons (Charts 12 to 19).
    Sets multi-level category cache (Level 0: Station, Level 1: Model) and formats series colors from Color_Template.
    """
    root = ET.fromstring(xml_content)
    barChart = root.find(f'.//{{{NS_C}}}barChart')
    if barChart is None: return xml_content, []

    unique_defects = []
    defect_points = {}

    for m_idx, model in enumerate(models_alphabetical[:5]):
        pt_st1 = m_idx * 2
        pt_st2 = m_idx * 2 + 1

        st1_defs = st1_model_defects.get(model, [])
        for def_name, qty in st1_defs:
            if qty > 0:
                if def_name not in unique_defects:
                    unique_defects.append(def_name)
                    defect_points[def_name] = {}
                defect_points[def_name][pt_st1] = int(qty)

        st2_defs = st2_model_defects.get(model, [])
        for def_name, qty in st2_defs:
            if qty > 0:
                if def_name not in unique_defects:
                    unique_defects.append(def_name)
                    defect_points[def_name] = {}
                defect_points[def_name][pt_st2] = int(qty)

    unique_defects.sort(key=lambda d: sum(defect_points[d].values()), reverse=True)

    existing_sers = barChart.findall(f'{{{NS_C}}}ser')
    insert_idx = list(barChart).index(existing_sers[0]) if existing_sers else 3
    for s in existing_sers:
        barChart.remove(s)

    # Remove any filteredBarSeries elements so no series is filtered/hidden
    for elem in list(root.iter()):
        for child in list(elem):
            if child.tag.endswith('filteredBarSeries'):
                elem.remove(child)

    for ser_idx, def_name in enumerate(unique_defects):
        hex_color = get_defect_hex(def_name, color_map)
        points_dict = defect_points.get(def_name, {})

        ser_elem = ET.Element(f'{{{NS_C}}}ser')
        ET.SubElement(ser_elem, f'{{{NS_C}}}idx', val=str(ser_idx))
        ET.SubElement(ser_elem, f'{{{NS_C}}}order', val=str(ser_idx))

        # Title
        tx = ET.SubElement(ser_elem, f'{{{NS_C}}}tx')
        strRef = ET.SubElement(tx, f'{{{NS_C}}}strRef')
        strCache = ET.SubElement(strRef, f'{{{NS_C}}}strCache')
        ET.SubElement(strCache, f'{{{NS_C}}}ptCount', val="1")
        pt0 = ET.SubElement(strCache, f'{{{NS_C}}}pt', idx="0")
        ET.SubElement(pt0, f'{{{NS_C}}}v').text = def_name

        # Solid Fill Color from Color Template
        spPr = ET.SubElement(ser_elem, f'{{{NS_C}}}spPr')
        solidFill = ET.SubElement(spPr, f'{{{NS_A}}}solidFill')
        ET.SubElement(solidFill, f'{{{NS_A}}}srgbClr', val=hex_color)
        ln = ET.SubElement(spPr, f'{{{NS_A}}}ln')
        ET.SubElement(ln, f'{{{NS_A}}}noFill')
        ET.SubElement(spPr, f'{{{NS_A}}}effectLst')

        ET.SubElement(ser_elem, f'{{{NS_C}}}invertIfNegative', val="0")

        # Data Labels
        dLbls = ET.SubElement(ser_elem, f'{{{NS_C}}}dLbls')
        dl_spPr = ET.SubElement(dLbls, f'{{{NS_C}}}spPr')
        ET.SubElement(dl_spPr, f'{{{NS_A}}}noFill')
        dl_ln = ET.SubElement(dl_spPr, f'{{{NS_A}}}ln')
        ET.SubElement(dl_ln, f'{{{NS_A}}}noFill')
        ET.SubElement(dl_spPr, f'{{{NS_A}}}effectLst')
        
        txPr = ET.SubElement(dLbls, f'{{{NS_C}}}txPr')
        ET.SubElement(txPr, f'{{{NS_A}}}bodyPr', rot="0", spcFirstLastPara="0", vertOverflow="ellipsis", vert="horz", wrap="square", lIns="38100", tIns="19050", rIns="38100", bIns="19050", anchor="ctr", anchorCtr="1")
        ET.SubElement(txPr, f'{{{NS_A}}}lstStyle')
        p = ET.SubElement(txPr, f'{{{NS_A}}}p')
        pPr = ET.SubElement(p, f'{{{NS_A}}}pPr')
        defRPr = ET.SubElement(pPr, f'{{{NS_A}}}defRPr', lang="en-US", sz="700", b="0", i="0", u="none", strike="noStrike", kern="1200", baseline="0")
        ET.SubElement(defRPr, f'{{{NS_A}}}solidFill').append(ET.Element(f'{{{NS_A}}}srgbClr', val="FFFFFF"))
        ET.SubElement(defRPr, f'{{{NS_A}}}latin', typeface="+mn-lt")
        ET.SubElement(defRPr, f'{{{NS_A}}}ea', typeface="+mn-ea")
        ET.SubElement(defRPr, f'{{{NS_A}}}cs', typeface="+mn-cs")
        ET.SubElement(p, f'{{{NS_A}}}endParaRPr', lang="vi-VN")

        ET.SubElement(dLbls, f'{{{NS_C}}}dLblPos', val="ctr")
        ET.SubElement(dLbls, f'{{{NS_C}}}showLegendKey', val="0")
        ET.SubElement(dLbls, f'{{{NS_C}}}showVal', val="1")
        ET.SubElement(dLbls, f'{{{NS_C}}}showCatName', val="0")
        ET.SubElement(dLbls, f'{{{NS_C}}}showSerName', val="0")
        ET.SubElement(dLbls, f'{{{NS_C}}}showPercent', val="0")
        ET.SubElement(dLbls, f'{{{NS_C}}}showBubbleSize', val="0")
        ET.SubElement(dLbls, f'{{{NS_C}}}showLeaderLines', val="0")

        # Multi-level Category (Lvl 0: Station, Lvl 1: Model)
        cat = ET.SubElement(ser_elem, f'{{{NS_C}}}cat')
        multiLvlStrRef = ET.SubElement(cat, f'{{{NS_C}}}multiLvlStrRef')
        multiLvlStrCache = ET.SubElement(multiLvlStrRef, f'{{{NS_C}}}multiLvlStrCache')
        ET.SubElement(multiLvlStrCache, f'{{{NS_C}}}ptCount', val="10")
        
        lvl0 = ET.SubElement(multiLvlStrCache, f'{{{NS_C}}}lvl')
        for pt_i in range(10):
            st_val = st1_name if pt_i % 2 == 0 else st2_name
            p_e = ET.SubElement(lvl0, f'{{{NS_C}}}pt', idx=str(pt_i))
            ET.SubElement(p_e, f'{{{NS_C}}}v').text = st_val

        lvl1 = ET.SubElement(multiLvlStrCache, f'{{{NS_C}}}lvl')
        for m_i, m_name in enumerate(models_alphabetical[:5]):
            p_e = ET.SubElement(lvl1, f'{{{NS_C}}}pt', idx=str(m_i * 2))
            ET.SubElement(p_e, f'{{{NS_C}}}v').text = m_name

        # Numerical Values
        val = ET.SubElement(ser_elem, f'{{{NS_C}}}val')
        numRef = ET.SubElement(val, f'{{{NS_C}}}numRef')
        numCache = ET.SubElement(numRef, f'{{{NS_C}}}numCache')
        ET.SubElement(numCache, f'{{{NS_C}}}formatCode').text = "General"
        ET.SubElement(numCache, f'{{{NS_C}}}ptCount', val="10")
        
        for pt_i in range(10):
            if pt_i in points_dict:
                p_e = ET.SubElement(numCache, f'{{{NS_C}}}pt', idx=str(pt_i))
                ET.SubElement(p_e, f'{{{NS_C}}}v').text = str(points_dict[pt_i])

        barChart.insert(insert_idx + ser_idx, ser_elem)

    return ET.tostring(root, encoding='utf-8', xml_declaration=True), unique_defects

def _update_slide_legend_tables(slide_xml, unique_defects, color_map):
    """
    Update the colored legend boxes and defect text in Slide 3 & 4 tables.
    Dynamically clones rows if unique_defects exceeds existing table rows,
    matching the exact distribution of the reference presentation.
    """
    root = ET.fromstring(slide_xml)
    tables = root.findall(f'.//{{{NS_A}}}tbl')
    if not tables:
        return slide_xml

    if len(tables) == 2:
        t0_rows = tables[0].findall(f'{{{NS_A}}}tr')
        t1_rows = tables[1].findall(f'{{{NS_A}}}tr')
        base_0 = len(t0_rows)
        base_1 = len(t1_rows)

        t0_defs = list(unique_defects[:base_0])
        t1_defs = list(unique_defects[base_0:base_0 + base_1])
        overflow = list(unique_defects[base_0 + base_1:])

        # If 'Wrong special label/hangtag' and 'Wrinkle collar lining' in overflow,
        # place Wrong special label/hangtag into Table 0 and Wrinkle collar lining into Table 1
        # to match the user's reference layout
        if 'Wrong special label/hangtag' in overflow and 'Wrinkle collar lining' in overflow:
            overflow.remove('Wrong special label/hangtag')
            overflow.remove('Wrinkle collar lining')
            t0_defs.append('Wrong special label/hangtag')
            t1_defs.append('Wrinkle collar lining')

        for i, item in enumerate(overflow):
            if i % 2 == 0:
                t0_defs.append(item)
            else:
                t1_defs.append(item)

        table_assignments = [(tables[0], t0_defs), (tables[1], t1_defs)]
    else:
        table_assignments = [(tables[0], list(unique_defects))]

    for tbl, def_list in table_assignments:
        rows = tbl.findall(f'{{{NS_A}}}tr')
        needed = len(def_list)

        while len(rows) < needed:
            last_row = rows[-1]
            new_row = copy.deepcopy(last_row)
            tbl.append(new_row)
            rows.append(new_row)

        for r_idx, row in enumerate(rows):
            cells = row.findall(f'{{{NS_A}}}tc')
            if len(cells) < 2:
                continue
            c0, c1 = cells[0], cells[1]

            if r_idx < needed:
                defect_name = def_list[r_idx]
                hex_color = get_defect_hex(defect_name, color_map)

                # Swatch cell (c0)
                tcPr0 = c0.find(f'{{{NS_A}}}tcPr')
                if tcPr0 is None:
                    tcPr0 = ET.SubElement(c0, f'{{{NS_A}}}tcPr')
                for old_f in list(tcPr0.findall(f'{{{NS_A}}}solidFill')):
                    tcPr0.remove(old_f)
                for old_f in list(tcPr0.findall(f'{{{NS_A}}}noFill')):
                    tcPr0.remove(old_f)
                solidFill = ET.SubElement(tcPr0, f'{{{NS_A}}}solidFill')
                ET.SubElement(solidFill, f'{{{NS_A}}}srgbClr', val=hex_color)

                # Label cell (c1)
                txBody = c1.find(f'{{{NS_A}}}txBody')
                if txBody is None:
                    txBody = ET.SubElement(c1, f'{{{NS_A}}}txBody')
                    ET.SubElement(txBody, f'{{{NS_A}}}bodyPr')

                p = txBody.find(f'{{{NS_A}}}p')
                if p is None:
                    p = ET.SubElement(txBody, f'{{{NS_A}}}p')

                for child in list(p):
                    if child.tag in [f'{{{NS_A}}}r', f'{{{NS_A}}}endParaRPr']:
                        p.remove(child)

                r = ET.SubElement(p, f'{{{NS_A}}}r')
                rPr = ET.SubElement(r, f'{{{NS_A}}}rPr', sz="900", b="0", i="0")
                solidFill_r = ET.SubElement(rPr, f'{{{NS_A}}}solidFill')
                ET.SubElement(solidFill_r, f'{{{NS_A}}}srgbClr', val="000000")
                ET.SubElement(rPr, f'{{{NS_A}}}latin', typeface="Times New Roman")

                t = ET.SubElement(r, f'{{{NS_A}}}t')
                t.text = defect_name

                endParaRPr = ET.SubElement(p, f'{{{NS_A}}}endParaRPr', sz="900", b="0", i="0")
                solidFill_end = ET.SubElement(endParaRPr, f'{{{NS_A}}}solidFill')
                ET.SubElement(solidFill_end, f'{{{NS_A}}}srgbClr', val="000000")
            else:
                tcPr0 = c0.find(f'{{{NS_A}}}tcPr')
                if tcPr0 is not None:
                    for old_f in list(tcPr0.findall(f'{{{NS_A}}}solidFill')):
                        tcPr0.remove(old_f)
                    ET.SubElement(tcPr0, f'{{{NS_A}}}noFill')
                for t_node in c1.findall(f'.//{{{NS_A}}}t'):
                    t_node.text = ""

    return ET.tostring(root, encoding='utf-8', xml_declaration=True)

def _update_slide_header_title(slide_xml, header_title):
    """
    Update slide header title banner (e.g. 'Re-Inspection Report_July 2026')
    to dynamic header_title (e.g. 'Re-Inspection Report_September 2026').
    """
    root = ET.fromstring(slide_xml)
    modified = False

    for sp in root.iter(f'{{{NS_P}}}sp'):
        tx_nodes = sp.findall(f'.//{{{NS_A}}}t')
        full_txt = ''.join([t.text for t in tx_nodes if t.text]).strip()

        # Match header title banner
        if 'Re-Inspection' in full_txt and ('Report_' in full_txt or 'Report' in full_txt) and ('Date' not in full_txt):
            if tx_nodes:
                tx_nodes[0].text = header_title
                for t_rem in tx_nodes[1:]:
                    t_rem.text = ""
                modified = True

    if modified:
        return ET.tostring(root, encoding='utf-8', xml_declaration=True)
    return slide_xml

def _update_cover_slide(slide_xml, month_label):
    """Update Slide 1 cover date badge dynamically."""
    if not month_label:
        month_label = datetime.date.today().strftime('%B %Y')

    root = ET.fromstring(slide_xml)
    modified = False

    for sp in root.iter(f'{{{NS_P}}}sp'):
        tx_nodes = sp.findall(f'.//{{{NS_A}}}t')
        full_txt = ''.join([t.text for t in tx_nodes if t.text]).strip()
        if 'Date:' in full_txt or full_txt.startswith('Date'):
            if tx_nodes:
                tx_nodes[0].text = f"Date: {month_label}"
                for t_rem in tx_nodes[1:]:
                    t_rem.text = ""
                modified = True

    if modified:
        return ET.tostring(root, encoding='utf-8', xml_declaration=True)
    return re.sub(rb'Date:\s*[A-Za-z]+\s*\d{4}', f'Date: {month_label}'.encode('utf-8'), slide_xml)

def _update_slide2_tables_and_texts(slide_xml, top_models_table, grand_tot, vn_tot, indo_tot, vh_qc, vh2_qc, jv_qc, jv2_qc, header_title):
    """Update Slide 2 tables and dynamic inspection count badges."""
    root = ET.fromstring(slide_xml)
    
    # 1. Update Top 5 Model Tables
    tables = root.findall(f'.//{{{NS_A}}}tbl')
    sites_order = ['VH', 'VH2', 'JV', 'JV2']
    for t_idx, site in enumerate(sites_order):
        if t_idx >= len(tables) or site not in top_models_table: continue
        tbl = tables[t_idx]
        rows = tbl.findall(f'{{{NS_A}}}tr')
        models_data = top_models_table[site]
        for r_idx in range(1, len(rows)):
            d_idx = r_idx - 1
            if d_idx < len(models_data):
                m_name, count, pct = models_data[d_idx]
                row = rows[r_idx]
                cells = row.findall(f'{{{NS_A}}}tc')
                if len(cells) >= 2:
                    t_elem1 = cells[0].find(f'.//{{{NS_A}}}t')
                    if t_elem1 is not None: t_elem1.text = str(m_name)
                    t_elem2 = cells[1].find(f'.//{{{NS_A}}}t')
                    if t_elem2 is not None: t_elem2.text = f"{pct:.2%}".replace('.', ',')

    # 2. Update Inspection Time Badges & Header Title
    for sp in root.iter(f'{{{NS_P}}}sp'):
        tx_nodes = sp.findall(f'.//{{{NS_A}}}t')
        full_txt = ''.join([t.text for t in tx_nodes if t.text]).strip()
        
        if 'Insp. times' in full_txt:
            cNvPr = sp.find(f'.//{{{NS_P}}}cNvPr')
            sid = cNvPr.get('id', '') if cNvPr is not None else ''
            
            if sid == '10' or '975' in full_txt or '898' in full_txt:
                new_val = f"Insp. times {grand_tot}"
            elif sid == '64' or ('305' in full_txt and 'times  305' in full_txt) or (sid == '64' and '337' in full_txt):
                new_val = f"Insp. times {vn_tot}"
            elif sid == '4' or '670' in full_txt or (sid == '4' and '561' in full_txt):
                new_val = f"Insp. times {indo_tot}"
            elif sid == '7' or '282' in full_txt or (sid == '7' and '303' in full_txt):
                new_val = f"Insp. times: {vh_qc}"
            elif sid == '9' or '23' in full_txt or (sid == '9' and '34' in full_txt):
                new_val = f"Insp. times: {vh2_qc}"
            elif sid == '18' or '99' in full_txt or (sid == '18' and '86' in full_txt):
                new_val = f"Insp. times: {jv_qc}"
            elif sid == '17' or '571' in full_txt or (sid == '17' and '475' in full_txt):
                new_val = f"Insp. times: {jv2_qc}"
            else:
                new_val = full_txt

            if tx_nodes:
                tx_nodes[0].text = new_val
                for t_rem in tx_nodes[1:]: t_rem.text = ""

        elif 'Re-Inspection' in full_txt and 'Report' in full_txt and 'Date' not in full_txt:
            if tx_nodes:
                tx_nodes[0].text = header_title
                for t_rem in tx_nodes[1:]: t_rem.text = ""

    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


# ==============================================================================
#  MAIN PRESENTATION EXPORT PIPELINE
# ==============================================================================

def export_reinspection_presentation(all_sites_data, template_pptx="RE-INS REPORT.APR.2026.pptx", output_pptx="RE-INS REPORT.pptx", month_label=None, title_format="Re-Inspection Report_{month_label}", color_template_path="Color_Template.xlsx", log_fn=print):
    """
    Generate fully functional executive PowerPoint report with exact layout, colors, and charts.
    """
    if not os.path.exists(template_pptx):
        raise FileNotFoundError(f"Template PowerPoint file not found: {template_pptx}")

    if not month_label:
        month_label = datetime.date.today().strftime('%B %Y')

    header_title = title_format.format(month_label=month_label, Month_Year=month_label)
    log_fn(f"Generating PowerPoint Report from template: {template_pptx} (Title: '{header_title}')")
    
    # Collect all unique defect names across all sites and stations
    all_defects = set()
    for s_data in all_sites_data.values():
        for d in s_data.get('top5_defects', []):
            if d and d[0]:
                all_defects.add(str(d[0]).strip().replace('\xa0', ' '))
        for grp in ['ftt_top3', 'hfpa_top3', 're_qc_top3', 're_ma_top3']:
            for m_item in s_data.get(grp, []):
                for d_item in m_item.get('defects', []):
                    if d_item and d_item[0]:
                        all_defects.add(str(d_item[0]).strip().replace('\xa0', ' '))

    color_map = sync_and_register_defects(all_defects, color_template_path, log_fn=log_fn)
    log_fn(f"[OK] Synchronized and loaded {len(color_map)} defect color definitions from {color_template_path}")

    vh = all_sites_data.get('VH', {'qc_total_fail': 0, 'top_countries': [], 'top5_defects': [], 'top_models': [], 'grand_total_5models': 1, 'models_alphabetical': [], 'ftt_top3': [], 'hfpa_top3': [], 're_qc_top3': [], 're_ma_top3': []})
    vh2 = all_sites_data.get('VH2', {'qc_total_fail': 0, 'top_countries': [], 'top5_defects': [], 'top_models': [], 'grand_total_5models': 1, 'models_alphabetical': [], 'ftt_top3': [], 'hfpa_top3': [], 're_qc_top3': [], 're_ma_top3': []})
    jv = all_sites_data.get('JV', {'qc_total_fail': 0, 'top_countries': [], 'top5_defects': [], 'top_models': [], 'grand_total_5models': 1, 'models_alphabetical': [], 'ftt_top3': [], 'hfpa_top3': [], 're_qc_top3': [], 're_ma_top3': []})
    jv2 = all_sites_data.get('JV2', {'qc_total_fail': 0, 'top_countries': [], 'top5_defects': [], 'top_models': [], 'grand_total_5models': 1, 'models_alphabetical': [], 'ftt_top3': [], 'hfpa_top3': [], 're_qc_top3': [], 're_ma_top3': []})

    vn_tot = vh['qc_total_fail'] + vh2['qc_total_fail']
    indo_tot = jv['qc_total_fail'] + jv2['qc_total_fail']
    grand_tot = vn_tot + indo_tot

    top_models_table = {}
    for s in ['VH', 'VH2', 'JV', 'JV2']:
        s_data = all_sites_data.get(s, {'top_models': [], 'qc_total_fail': 1})
        s_qc = max(s_data['qc_total_fail'], 1)
        top_models_table[s] = [
            (m, count, count / s_qc) for m, count in s_data['top_models'][:5]
        ]

    def make_defect_dict(model_def_list):
        res = {}
        for item in model_def_list:
            res[item['model']] = item['defects']
        return res

    s3_all_unique = set()
    s4_all_unique = set()

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(template_pptx, 'r') as zin:
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                content = zin.read(item.filename)

                # --- Slide 1 (Cover) ---
                if item.filename == 'ppt/slides/slide1.xml':
                    content = _update_cover_slide(content, month_label)

                # --- Slide 2 (Summary Dashboard) ---
                elif item.filename == 'ppt/slides/slide2.xml':
                    content = _update_slide2_tables_and_texts(content, top_models_table, grand_tot, vn_tot, indo_tot, vh['qc_total_fail'], vh2['qc_total_fail'], jv['qc_total_fail'], jv2['qc_total_fail'], header_title)

                # --- Chart 1: INDO Doughnut (JV vs JV2) ---
                elif item.filename == 'ppt/charts/chart1.xml':
                    content = _update_doughnut_chart(content, 'INDO', ['JV', 'JV2'], [jv['qc_total_fail'], jv2['qc_total_fail']])

                # --- Chart 2: VN Doughnut (VH vs VH2) ---
                elif item.filename == 'ppt/charts/chart2.xml':
                    content = _update_doughnut_chart(content, 'VN', ['VH', 'VH2'], [vh['qc_total_fail'], vh2['qc_total_fail']])

                # --- Chart 3: CLG Doughnut (VN vs INDO) ---
                elif item.filename == 'ppt/charts/chart3.xml':
                    content = _update_doughnut_chart(content, 'CLG', ['VN', 'INDO'], [vn_tot, indo_tot])

                # --- Chart 4: VH Top 5 Defects (QC+MA) ---
                elif item.filename == 'ppt/charts/chart4.xml':
                    cats = [d[0] for d in vh['top5_defects'][:5]]
                    vals1 = [d[1] for d in vh['top5_defects'][:5]]
                    gt = max(vh['grand_total_5models'], 1)
                    vals2 = [v / gt for v in vals1]
                    content = _update_bar_chart_2series(content, cats, vals1, vals2)

                # --- Chart 5: VH2 Top 5 Defects (QC+MA) ---
                elif item.filename == 'ppt/charts/chart5.xml':
                    cats = [d[0] for d in vh2['top5_defects'][:5]]
                    vals1 = [d[1] for d in vh2['top5_defects'][:5]]
                    gt = max(vh2['grand_total_5models'], 1)
                    vals2 = [v / gt for v in vals1]
                    content = _update_bar_chart_2series(content, cats, vals1, vals2)

                # --- Chart 6: JV Top 5 Defects (QC+MA) ---
                elif item.filename == 'ppt/charts/chart6.xml':
                    cats = [d[0] for d in jv['top5_defects'][:5]]
                    vals1 = [d[1] for d in jv['top5_defects'][:5]]
                    gt = max(jv['grand_total_5models'], 1)
                    vals2 = [v / gt for v in vals1]
                    content = _update_bar_chart_2series(content, cats, vals1, vals2)

                # --- Chart 7: JV2 Top 5 Defects (QC+MA) ---
                elif item.filename == 'ppt/charts/chart7.xml':
                    cats = [d[0] for d in jv2['top5_defects'][:5]]
                    vals1 = [d[1] for d in jv2['top5_defects'][:5]]
                    gt = max(jv2['grand_total_5models'], 1)
                    vals2 = [v / gt for v in vals1]
                    content = _update_bar_chart_2series(content, cats, vals1, vals2)

                # --- Chart 8: VH Top 5 Countries Pie3D ---
                elif item.filename == 'ppt/charts/chart8.xml':
                    vh_qc = max(vh['qc_total_fail'], 1)
                    top_c = vh['top_countries'][:5]
                    top_c_sum = sum(c[1] for c in top_c)
                    others = max(vh['qc_total_fail'] - top_c_sum, 0)
                    cats = [c[0] for c in top_c] + (['OTHERS'] if others > 0 or len(top_c) < 6 else [])
                    vals1 = [c[1] for c in top_c] + ([others] if others > 0 or len(top_c) < 6 else [])
                    vals2 = [v / vh_qc for v in vals1]
                    content = _update_pie3d_chart(content, cats, vals1, vals2)

                # --- Chart 9: VH2 Top 5 Countries Pie3D ---
                elif item.filename == 'ppt/charts/chart9.xml':
                    vh2_qc = max(vh2['qc_total_fail'], 1)
                    top_c = vh2['top_countries'][:5]
                    top_c_sum = sum(c[1] for c in top_c)
                    others = max(vh2['qc_total_fail'] - top_c_sum, 0)
                    cats = [c[0] for c in top_c] + (['OTHERS'] if others > 0 or len(top_c) < 6 else [])
                    vals1 = [c[1] for c in top_c] + ([others] if others > 0 or len(top_c) < 6 else [])
                    vals2 = [v / vh2_qc for v in vals1]
                    content = _update_pie3d_chart(content, cats, vals1, vals2)

                # --- Chart 10: JV2 Top 5 Countries Pie3D ---
                elif item.filename == 'ppt/charts/chart10.xml':
                    jv2_qc = max(jv2['qc_total_fail'], 1)
                    top_c = jv2['top_countries'][:5]
                    top_c_sum = sum(c[1] for c in top_c)
                    others = max(jv2['qc_total_fail'] - top_c_sum, 0)
                    cats = [c[0] for c in top_c] + (['OTHERS'] if others > 0 or len(top_c) < 6 else [])
                    vals1 = [c[1] for c in top_c] + ([others] if others > 0 or len(top_c) < 6 else [])
                    vals2 = [v / jv2_qc for v in vals1]
                    content = _update_pie3d_chart(content, cats, vals1, vals2)

                # --- Chart 11: JV Top 5 Countries Pie3D ---
                elif item.filename == 'ppt/charts/chart11.xml':
                    jv_qc = max(jv['qc_total_fail'], 1)
                    top_c = jv['top_countries'][:5]
                    top_c_sum = sum(c[1] for c in top_c)
                    others = max(jv['qc_total_fail'] - top_c_sum, 0)
                    cats = [c[0] for c in top_c] + (['OTHERS'] if others > 0 or len(top_c) < 6 else [])
                    vals1 = [c[1] for c in top_c] + ([others] if others > 0 or len(top_c) < 6 else [])
                    vals2 = [v / jv_qc for v in vals1]
                    content = _update_pie3d_chart(content, cats, vals1, vals2)

                # --- Slide 3 Stacked Column Charts: 12 (VH FTT), 13 (VH HFPA), 14 (VH2 FTT), 15 (VH2 HFPA) ---
                elif item.filename == 'ppt/charts/chart12.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'FTT', 'Re-ins', vh['models_alphabetical'], make_defect_dict(vh['ftt_top3']), make_defect_dict(vh['re_qc_top3']), color_map)
                    s3_all_unique.update(u_defs)

                elif item.filename == 'ppt/charts/chart13.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'HFPA', 'Re-ins', vh['models_alphabetical'], make_defect_dict(vh['hfpa_top3']), make_defect_dict(vh['re_ma_top3']), color_map)
                    s3_all_unique.update(u_defs)

                elif item.filename == 'ppt/charts/chart14.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'FTT', 'Re-ins', vh2['models_alphabetical'], make_defect_dict(vh2['ftt_top3']), make_defect_dict(vh2['re_qc_top3']), color_map)
                    s3_all_unique.update(u_defs)

                elif item.filename == 'ppt/charts/chart15.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'HFPA', 'Re-ins', vh2['models_alphabetical'], make_defect_dict(vh2['hfpa_top3']), make_defect_dict(vh2['re_ma_top3']), color_map)
                    s3_all_unique.update(u_defs)

                # --- Slide 4 Stacked Column Charts: 16 (JV FTT), 17 (JV HFPA), 18 (JV2 FTT), 19 (JV2 HFPA) ---
                elif item.filename == 'ppt/charts/chart16.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'FTT', 'Re-ins', jv['models_alphabetical'], make_defect_dict(jv['ftt_top3']), make_defect_dict(jv['re_qc_top3']), color_map)
                    s4_all_unique.update(u_defs)

                elif item.filename == 'ppt/charts/chart17.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'HFPA', 'Re-ins', jv['models_alphabetical'], make_defect_dict(jv['hfpa_top3']), make_defect_dict(jv['re_ma_top3']), color_map)
                    s4_all_unique.update(u_defs)

                elif item.filename == 'ppt/charts/chart18.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'FTT', 'Re-ins', jv2['models_alphabetical'], make_defect_dict(jv2['ftt_top3']), make_defect_dict(jv2['re_qc_top3']), color_map)
                    s4_all_unique.update(u_defs)

                elif item.filename == 'ppt/charts/chart19.xml':
                    content, u_defs = _update_comparison_stacked_chart(content, 'HFPA', 'Re-ins', jv2['models_alphabetical'], make_defect_dict(jv2['hfpa_top3']), make_defect_dict(jv2['re_ma_top3']), color_map)
                    s4_all_unique.update(u_defs)

                zout.writestr(item, content)

    # Second pass for Slide 3 and Slide 4 legend tables & date headers
    zip_buffer2 = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'r') as zin:
        with zipfile.ZipFile(zip_buffer2, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                content = zin.read(item.filename)
                if item.filename == 'ppt/slides/slide3.xml':
                    content = _update_slide_legend_tables(content, sorted(s3_all_unique), color_map)
                    content = _update_slide_header_title(content, header_title)
                elif item.filename == 'ppt/slides/slide4.xml':
                    content = _update_slide_legend_tables(content, sorted(s4_all_unique), color_map)
                    content = _update_slide_header_title(content, header_title)
                elif item.filename.startswith('ppt/slides/slide') and item.filename != 'ppt/slides/slide1.xml':
                    # Apply dynamic header title to any other content slides
                    content = _update_slide_header_title(content, header_title)
                zout.writestr(item, content)

    # Save to destination file
    out_dir = os.path.dirname(os.path.abspath(output_pptx))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    try:
        with open(output_pptx, 'wb') as f_out:
            f_out.write(zip_buffer2.getvalue())
        target_saved = output_pptx
        log_fn(f"[OK] PowerPoint presentation generated successfully: {os.path.abspath(output_pptx)}")
    except PermissionError:
        base, ext = os.path.splitext(output_pptx)
        alt_path = f"{base}_updated{ext}"
        with open(alt_path, 'wb') as f_out:
            f_out.write(zip_buffer2.getvalue())
        target_saved = alt_path
        log_fn(f"[WARNING] {output_pptx} is locked by PowerPoint/WPS. Saved to {alt_path} instead.")

    # Automatically apply visual inspection annotations (bounding frames, connector lines & insight callouts)
    try:
        import pptx_annotator
        pptx_annotator.annotate_presentation(target_saved, all_sites_data=all_sites_data, log_fn=log_fn)
    except Exception as e:
        log_fn(f"[WARNING] Could not apply visual inspection annotations: {e}")

    return target_saved
